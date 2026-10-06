"""Stock tools (5 Oct §4.8 / §4.14, T-09).

* Stock Reconciliation - scan barcodes (or import the scanner's TXT) against
  what the books say is in stock: three lists - in stock and scanned (OK),
  in stock but not scanned (missing), scanned but not in stock (unknown,
  sold, out on approval, at another location).
* Ready Closing Stock - every piece in stock with its classification,
  weights, fine, the day's metal rate, price and tag.
* SKU Status - every piece ever made of one SKU, with where it is now.
* Stock View - the stock browser with status / location / stone / other
  filters.
* Barcode Catalog - pieces by Stock ID, Job No or SKU for a catalog or tags.
"""
from __future__ import annotations

import json
import re
import shutil
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from diagold.db.models import (Account, FamilyCategory, Job, Location, Metal, MfgTransfer,
                               MfgTransferLine, ProductSku, StockItem, StockRecon,
                               StockReconScan)
from diagold.db.models.sku import Item
from diagold.services import costing, production

ZERO = Decimal("0")
D3 = Decimal("0.001")
STATUS_TEXT = {"in_stock": "In-Stock", "sold": "Sold", "on_approval": "Approval",
               "in_repair": "Repair", "returned": "Returned", "melted": "Melted"}
VIEW_STATUSES = {"In Stock": ("in_stock",), "Approval": ("on_approval",), "Sold": ("sold",),
                 "All Stock": tuple(STATUS_TEXT)}


def _dec(v: Any) -> Decimal:
    if v in (None, ""):
        return ZERO
    return v if isinstance(v, Decimal) else Decimal(str(v))


class _Cache:
    def __init__(self, session: Session):
        self.s = session
        self._d: dict[tuple, Any] = {}

    def get(self, model, pk):
        if not pk:
            return None
        key = (model, pk)
        if key not in self._d:
            self._d[key] = self.s.get(model, pk)
        return self._d[key]


def stones_of(session: Session, item: StockItem) -> list[dict]:
    if item.line_id:
        line = session.get(MfgTransferLine, item.line_id)
        if line is not None:
            return json.loads(line.stones_json or "[]")
    if item.in_line_id:
        from diagold.db.models import ReadyVoucherLine
        src = session.get(ReadyVoucherLine, item.in_line_id)
        if src is not None:
            return json.loads(src.stones_json or "[]")
    return []


def piece_row(session: Session, item: StockItem, c: _Cache | None = None,
              on_date: date | None = None) -> dict[str, Any]:
    """One piece, every column the stock tools show."""
    c = c or _Cache(session)
    job = c.get(Job, item.job_id)
    sku = c.get(ProductSku, item.product_sku_id)
    fam = c.get(FamilyCategory, sku.family_id) if sku else None
    itm = c.get(Item, sku.item_id) if sku else None
    metal_id = item.metal_id or (job.metal_id if job else None)
    metal = c.get(Metal, metal_id)
    loc = c.get(Location, item.location_id)
    line = c.get(MfgTransferLine, item.line_id)
    tr = c.get(MfgTransfer, line.transfer_id) if line else None
    client = c.get(Account, job.account_id) if job else None
    holder = c.get(Account, item.holder_account_id)
    net, gross = _dec(item.net_wt), _dec(item.gross_wt)
    stones = stones_of(session, item)
    st_wt = sum((_dec(s.get("weight")) for s in stones), ZERO)
    rate = production.metal_price(session, metal_id, on_date or date.today())
    pcs = int(item.pcs or 1)
    return {
        "_id": item.id, "_job_id": job.id if job else None, "_stones": stones,
        "barcode": item.stock_no, "stock_no": item.stock_no,
        "location": loc.name if loc else "", "job_no": job.job_no if job else "",
        "sku": sku.sku_code if sku else "", "cf": client.code if client else "",
        "client": client.name if client else "",
        "category": sku.category if sku else "", "sub_category": sku.category2 if sku else "",
        "family": fam.name if fam else "", "item": itm.name if itm else "",
        "style": sku.style if sku else "", "pattern": sku.pattern if sku else "",
        "pattern2": sku.pattern2 if sku else "", "sku_ref": sku.sku_ref if sku else "",
        "master_sku": sku.master_sku if sku else "",
        "_cad": bool(sku.is_cad) if sku else False,
        "_master": bool(sku.is_master) if sku else False,
        "metal": metal.name if metal else "", "col": item.colour or (job.colour if job else ""),
        "size": item.size or (production._order_line_size(session, job) if job else ""),
        "c_ref": item.c_ref or (job.c_ref if job else ""), "pcs": pcs,
        "g_wt": gross, "n_wt": net, "st_wt": st_wt.quantize(D3) if st_wt else None,
        "oth_wt": (gross - net - st_wt / 5).quantize(D3) if gross > net + st_wt / 5 else None,
        "fine": costing.fine_metal_weight(metal, net).quantize(D3) if metal else None,
        "mt_rate": rate or None, "price": (_dec(item.price) / pcs).quantize(Decimal("0.01")),
        "amount": _dec(item.price), "tag": item.tag_text,
        "status": STATUS_TEXT.get(item.status, item.status), "_status": item.status,
        "holder": holder.name if holder else "",
        "date": tr.vr_date if tr else None,
        "_photo": (sku.image_finished or sku.image_design) if sku else "",
    }


# --------------------------------------------------------------------------
# Ready Closing Stock
# --------------------------------------------------------------------------
def ready_closing_stock(session: Session, date_to: date) -> list[dict[str, Any]]:
    """Every piece in stock: LOCATION, JOB NO, SKU, CF, Category, Sub-Category,
    Family, Item, METAL, COL, SIZE, PCS, G / N / OTH WT, FINE, MT-RATE (the
    day's rate per gram of that metal), PRICE (per pc), AMOUNT, TAG.
    OTH-WT = gross - net - stones (ct / 5 as grams)."""
    c = _Cache(session)
    rows = []
    for item in session.scalars(select(StockItem).where(StockItem.status == "in_stock")
                                .order_by(StockItem.stock_no)):
        r = piece_row(session, item, c, date_to)
        if r["date"] is not None and r["date"] > date_to:
            continue
        rows.append(r)
    rows.sort(key=lambda r: (r["location"], r["sku"], r["stock_no"]))
    return rows


# --------------------------------------------------------------------------
# SKU Status
# --------------------------------------------------------------------------
def sku_status(session: Session, sku_text: str) -> list[dict[str, Any]]:
    """Every piece ever made (or bought) of one SKU, on its latest voucher,
    and the SKU's jobs still in manufacturing (STATUS = MFG)."""
    from diagold.services import sales
    sku_text = (sku_text or "").strip()
    if not sku_text:
        return []
    sku = session.scalar(select(ProductSku).where(
        func.lower(ProductSku.sku_code) == sku_text.lower()))
    if sku is None:
        return []
    c = _Cache(session)
    rows = []
    for item in session.scalars(select(StockItem).where(StockItem.product_sku_id == sku.id)
                                .order_by(StockItem.stock_no)):
        p = piece_row(session, item, c)
        hist = sales.piece_history(session, item)
        last = hist[-1] if hist else {}
        vrtype = last.get("vrtype", "")
        if not hist and item.source == "opening":
            vrtype = "OPR"
        rows.append({**p, "date": last.get("date"), "vrtype": vrtype,
                     "vrno": last.get("vrno", ""), "particulars": last.get("party", ""),
                     "refno": "", "location": last.get("location") or p["location"]})
    transferred = {r["_job_id"] for r in rows if r["_job_id"]}
    for job in session.scalars(select(Job).where(
            Job.product_sku_id == sku.id,
            Job.status.in_(("pending", "mapped", "in_progress", "complete")))
            .order_by(Job.job_no)):
        if job.id in transferred:
            continue
        metal = c.get(Metal, job.metal_id)
        client = c.get(Account, job.account_id)
        step, _issue = production.current_step(session, job)
        gross, net = production.last_weights(session, job)
        rows.append({"_id": None, "_job_id": job.id, "date": job.mapped_on or job.prod_date,
                     "vrtype": "MFG", "vrno": "", "particulars": client.name if client else "",
                     "refno": "", "location": "", "job_no": job.job_no, "barcode": "",
                     "c_ref": job.c_ref, "item": "", "metal": metal.name if metal else "",
                     "col": job.colour, "size": production._order_line_size(session, job),
                     "pcs": job.pcs, "g_wt": gross, "n_wt": net, "st_wt": None,
                     "price": None, "status": "MFG" + (" (pending MFG transfer)"
                                                       if job.status == "complete" else ""),
                     "_photo": sku.image_finished or sku.image_design, "sku": sku.sku_code})
    return rows


# --------------------------------------------------------------------------
# Stock View
# --------------------------------------------------------------------------
STONE_FILTERS = ("ssku", "stone_group", "shape", "stone_type", "quality", "size")
OTHER_FILTERS = ("c_ref", "pattern", "pattern2", "category", "sub_category", "style",
                 "sku_ref", "master_sku", "sku")


def _has(value: str, wanted: str) -> bool:
    return wanted.strip().casefold() in (value or "").casefold()


def stock_view(session: Session, *, status: str = "In Stock",
               location_ids: Iterable[int] | None = None,
               stone: dict[str, str] | None = None, other: dict[str, str] | None = None,
               date_from: date | None = None, date_to: date | None = None,
               cad: bool = False, master: bool = False, approval_no: str = "",
               on_date: date | None = None) -> list[dict[str, Any]]:
    """The Stock View browser: pieces by status and location, narrowed by
    stone fields (any of the piece's stones matching all of them) and other
    fields (text contains, ignoring case)."""
    from diagold.services import price_charts as PC
    statuses = VIEW_STATUSES.get(status, ("in_stock",))
    q = select(StockItem).where(StockItem.status.in_(statuses)).order_by(StockItem.stock_no)
    locs = set(location_ids or ())
    if locs:
        q = q.where(StockItem.location_id.in_(locs))
    stone = {k: v for k, v in (stone or {}).items() if (v or "").strip()}
    other = {k: v for k, v in (other or {}).items() if (v or "").strip()}
    on_approval: set[int] | None = None
    if approval_no.strip().isdigit():
        from diagold.db.models import ReadyVoucher, ReadyVoucherLine
        on_approval = set(session.scalars(
            select(ReadyVoucherLine.stock_item_id).join(ReadyVoucher).where(
                ReadyVoucher.vr_type == "rs_approval",
                ReadyVoucher.vr_no == int(approval_no))))
    c = _Cache(session)
    rows = []
    for item in session.scalars(q):
        if on_approval is not None and item.id not in on_approval:
            continue
        r = piece_row(session, item, c, on_date)
        if cad and not r["_cad"] or master and not r["_master"]:
            continue
        if date_from and (r["date"] is None or r["date"] < date_from):
            continue
        if date_to and r["date"] is not None and r["date"] > date_to:
            continue
        if any(not _has(str(r.get(k) or ""), v) for k, v in other.items()):
            continue
        if stone:
            ok = False
            for st in r["_stones"]:
                k = PC.stone_keys(session, st)
                k["stone_group"] = k.pop("group")
                if all(_has(str(k.get(f) or ""), v) for f, v in stone.items()):
                    ok = True
                    break
            if not ok:
                continue
        rows.append(r)
    return rows


# --------------------------------------------------------------------------
# Barcode Catalog: pieces by Stock ID / Job No / SKU
# --------------------------------------------------------------------------
def numbers_in(text: str) -> list[int]:
    return [int(n) for n in re.findall(r"\d+", text or "")]


def find_pieces(session: Session, by: str, values: Iterable[str]) -> list[StockItem]:
    """by = "stock" / "job" / "sku"; values as typed or read from a TXT file."""
    vals = [v.strip() for v in values if v and v.strip()]
    if not vals:
        return []
    if by == "stock":
        nos = [int(v) for v in vals if v.isdigit()]
        q = select(StockItem).where(StockItem.stock_no.in_(nos))
    elif by == "job":
        nos = [int(v) for v in vals if v.isdigit()]
        q = select(StockItem).join(Job, StockItem.job_id == Job.id).where(Job.job_no.in_(nos))
    else:
        low = [v.lower() for v in vals]
        q = select(StockItem).join(ProductSku, StockItem.product_sku_id == ProductSku.id).where(
            func.lower(ProductSku.sku_code).in_(low))
    return list(session.scalars(q.order_by(StockItem.stock_no)))


def catalog_html(rows: list[dict[str, Any]], title: str = "Catalog",
                 per_row: int = 3, small: bool = False) -> str:
    """Catalog: a photo card per piece. Catalog 4x8 = 4 across, small cards
    (32 to an A4 page)."""
    h = 70 if small else 120
    cells = []
    for i, r in enumerate(rows):
        img = (f"<img src='{r['_photo']}' height={h}><br>" if r.get("_photo")
               else f"<div style='height:{h}px'>[no photo]</div>")
        cells.append(
            f"<td width={100 // per_row}% valign=top align=center "
            "style='border:1px solid #ccc;padding:4px;font-size:"
            f"{'8' if small else '10'}pt'>{img}<b>{r.get('sku', '')}</b><br>"
            f"Stock {r.get('stock_no', '')} · Job {r.get('job_no', '')}<br>"
            f"G {_dec(r.get('g_wt')):.3f} · N {_dec(r.get('n_wt')):.3f}<br>"
            f"{r.get('metal', '')} {r.get('col', '')} · Tag {r.get('tag', '')}</td>")
        if (i + 1) % per_row == 0:
            cells.append("</tr><tr>")
    return (f"<h2>{title}</h2><p>{len(rows)} piece(s)</p>"
            f"<table width=100% cellspacing=4><tr>{''.join(cells)}</tr></table>")


def create_image_folder(rows: list[dict[str, Any]], folder: str) -> int:
    """Create Image Folder: copy each piece's photo, named by Stock No."""
    out = Path(folder)
    out.mkdir(parents=True, exist_ok=True)
    n = 0
    for r in rows:
        src = Path(r.get("_photo") or "")
        if src.is_file():
            shutil.copy2(src, out / f"{r['stock_no']}{src.suffix}")
            n += 1
    return n


# --------------------------------------------------------------------------
# Stock Reconciliation
# --------------------------------------------------------------------------
def new_recon(session: Session, on_date: date, location_id: int | None = None,
              remark: str = "", user_id: int | None = None) -> StockRecon:
    no = (session.scalar(select(func.max(StockRecon.recon_no))) or 0) + 1
    r = StockRecon(recon_no=no, recon_date=on_date, location_id=location_id, remark=remark,
                   user_id=user_id)
    session.add(r)
    session.flush()
    return r


def add_scans(session: Session, recon: StockRecon, codes: Iterable[str]) -> tuple[int, int]:
    """Scanned or imported barcodes; a repeat scan is counted once.
    Returns (added, repeats)."""
    have = {s.barcode for s in recon.scans}
    added = dup = 0
    for code in codes:
        code = (code or "").strip()
        if not code:
            continue
        if code in have:
            dup += 1
            continue
        have.add(code)
        recon.scans.append(StockReconScan(barcode=code))
        added += 1
    session.flush()
    return added, dup


def reconcile(session: Session, recon: StockRecon) -> dict[str, list[dict[str, Any]]]:
    """ok      = in stock (at the location, if one is set) and scanned;
    missing = in stock but not scanned ("Add Stock But Not Show In Stock");
    unknown = scanned but not in stock there ("Data Unfound But In
              Reconciliation") - with why: sold, on approval, another
              location, or no such barcode."""
    c = _Cache(session)
    q = select(StockItem).where(StockItem.status == "in_stock")
    if recon.location_id:
        q = q.where(StockItem.location_id == recon.location_id)
    stock = {str(i.stock_no): i for i in session.scalars(q)}
    scanned = [s.barcode for s in recon.scans]
    seen = set(scanned)
    ok = [piece_row(session, stock[b], c, recon.recon_date) for b in scanned if b in stock]
    missing = [piece_row(session, i, c, recon.recon_date)
               for b, i in sorted(stock.items(), key=lambda kv: int(kv[0]))
               if b not in seen]
    unknown = []
    for b in scanned:
        if b in stock:
            continue
        item = session.scalar(select(StockItem).where(StockItem.stock_no == int(b))) \
            if b.isdigit() else None
        if item is None:
            unknown.append({"barcode": b, "status": "Not found - no such barcode",
                            "_id": None, "_job_id": None})
            continue
        r = piece_row(session, item, c, recon.recon_date)
        if item.status == "in_stock":
            r["status"] = f"In stock at {r['location']} (another location)"
        elif r["holder"]:
            r["status"] += f" - {r['holder']}"
        unknown.append(r)
    return {"ok": ok, "missing": missing, "unknown": unknown}


# --------------------------------------------------------------------------
# Item Search helpers (5 Oct §4.1, T-16)
# --------------------------------------------------------------------------
def sku_made_left(session: Session, sku_id: int | None) -> tuple[int, int]:
    """"This SKU Stock made / left": pieces ever made (or bought) of the SKU,
    and how many are still in stock."""
    if not sku_id:
        return 0, 0
    made = session.scalar(select(func.count(StockItem.id)).where(
        StockItem.product_sku_id == sku_id)) or 0
    left = session.scalar(select(func.count(StockItem.id)).where(
        StockItem.product_sku_id == sku_id, StockItem.status == "in_stock")) or 0
    return made, left


def cert_xlsx(session: Session, item: StockItem, path: str) -> None:
    """Cert Excel: the piece's certificate data - identity, weights, metal,
    every stone line - for the certificate lab / the client."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    r = piece_row(session, item)
    wb = Workbook()
    ws = wb.active
    ws.title = "Certificate"
    bold = Font(bold=True)
    for k, label in (("stock_no", "Stock No"), ("sku", "SKU"), ("job_no", "Job No"),
                     ("item", "Item"), ("family", "Family"), ("metal", "Metal"),
                     ("col", "Colour"), ("size", "Size"), ("pcs", "Pcs"),
                     ("g_wt", "Gross Wt (g)"), ("n_wt", "Net Wt (g)"), ("fine", "Fine (g)"),
                     ("st_wt", "Stone Wt (ct)")):
        v = r.get(k)
        ws.append([label, float(v) if isinstance(v, Decimal) else v])
        ws.cell(ws.max_row, 1).font = bold
    ws.append(["Cert No", item.cert_no or ""])
    ws.append(["HUID", item.huid or ""])
    ws.append([])
    ws.append(["Stone", "S Type", "Size", "Pcs", "Weight (ct)"])
    for c in ws[ws.max_row]:
        c.font = bold
    for st in r["_stones"]:
        ws.append([st.get("label") or "", st.get("s_type") or "", st.get("size") or "",
                   int(st.get("pcs") or 0), float(_dec(st.get("weight")))])
    ws.column_dimensions["A"].width = 22
    wb.save(path)
