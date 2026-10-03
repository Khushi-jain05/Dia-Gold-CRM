"""Job Costing (2 Oct Session 3 R1 / T-02): "Job costing zaroori hai - kisi
job ko karne me kya costing ho rahi hai."

One sheet per job, as the legacy Job Costing Sheet reads it (BANG-52):

    Metal        = N-Wt x metal rate                 12.238 x 8,831.40 = 1,08,078.67
    Stones       = sum(ct or pcs x price)            40,670.00
    Labour       = Setting + STD                     4,059.00 + 14,685.60
    Total        = Metal + Stones + Labour           1,67,493.27
    Margin       = Total x Margin %                  50% = 83,746.64
    Grand Total  = Total + Margin                    2,51,239.91
    Tag price    = Grand Total / 1000 (tag rule)     251        (UNCONFIRMED - Q3)
    Price per gm = tag price / G-Wt                  19.66

Setting is the setting labour stored on the job's receipts (pieces set x
setting-type rate, 28 Sept T-05) - the legacy sheet shows it beside STD.

A job already in ready stock is costed from what its MFG transfer froze (metal
rate, stone prices, labour rate, margin) - never silently re-priced. Any other
job is costed at the masters as on the date asked; a job still in work (WIP
costing) uses its net weight as last weighed and the stones issued so far.
"""
from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from diagold.db.models import Job, JobVoucher, Metal, MfgTransferLine, StockItem
from diagold.services import mfg_pricing, production

ZERO = Decimal("0")
PAISA = Decimal("0.01")
FINISHED = ("complete", "transferred")
WIP = ("mapped", "in_progress")


def _dec(v: Any) -> Decimal:
    if v is None or v == "":
        return ZERO
    return v if isinstance(v, Decimal) else Decimal(str(v))


@dataclass
class CostingSheet:
    job_id: int
    job_no: int
    sku: str
    metal: str
    colour: str
    size: str
    pcs: int
    gross_wt: Decimal
    price: mfg_pricing.TransferPrice
    basis: str                          # where the rates came from
    stock_no: int | None = None
    photo: str = ""
    wip: bool = False
    notes: list[str] = field(default_factory=list)

    # the sheet's own lines, named as the legacy sheet names them
    @property
    def metal_amount(self) -> Decimal:
        return self.price.metal_amount

    @property
    def stone_pcs(self) -> int:
        return sum(s.pcs for s in self.price.stones)

    @property
    def stone_wt(self) -> Decimal:
        return sum((s.weight for s in self.price.stones), ZERO)

    @property
    def setting(self) -> Decimal:
        return self.price.setting_amount

    @property
    def std_labour(self) -> Decimal:
        return self.price.labour

    @property
    def labour(self) -> Decimal:
        """Setting + STD + findings / extra metal - the LABOUR column."""
        p = self.price
        return p.setting_amount + p.labour + p.finding_labour + p.ex_metal_amount

    @property
    def total(self) -> Decimal:
        return self.price.total

    @property
    def grand_total(self) -> Decimal:
        return self.price.price_per_pcs

    @property
    def tag_value(self) -> Decimal:
        try:
            return Decimal(self.price.tag_text)
        except Exception:  # noqa: BLE001 - a non-numeric tag rule
            return ZERO

    @property
    def price_per_gm(self) -> Decimal:
        """Tag price per gross gram, as the legacy sheet shows (251 / 12.766)."""
        if not self.gross_wt:
            return ZERO
        return (self.tag_value / self.gross_wt).quantize(PAISA, rounding=ROUND_HALF_UP)

    def stones_by_group(self) -> dict[str, tuple[int, Decimal, Decimal]]:
        """Stone Group Wise (Ctrl+F1): pcs, ct and amount per DIAMOND / POLKI /
        COLOR STONE …"""
        out: dict[str, list] = {}
        for s in self.price.stones:
            g = (s.s_type or "Other").strip() or "Other"
            acc = out.setdefault(g, [0, ZERO, ZERO])
            acc[0] += s.pcs
            acc[1] += s.weight
            acc[2] += s.amount
        return {k: (v[0], v[1], v[2]) for k, v in out.items()}


def setting_labour(session: Session, job: Job) -> Decimal:
    """Setting labour earned on the job - stored on its receipts (T-05)."""
    total = session.scalar(select(func.sum(JobVoucher.labour)).where(
        JobVoucher.job_id == job.id, JobVoucher.kind == "receive"))
    return _dec(total).quantize(PAISA)


def _transfer_line(session: Session, job: Job) -> MfgTransferLine | None:
    return session.scalars(select(MfgTransferLine).where(MfgTransferLine.job_id == job.id)
                           .order_by(MfgTransferLine.id.desc())).first()


def costing_sheet(session: Session, job: Job, on_date: date | None = None, *,
                  live: bool = False) -> CostingSheet:
    """The Job Costing Sheet for one job. ``live`` costs it at the masters as
    on ``on_date`` even if it was transferred (the sheet's WIP Costing)."""
    from diagold.services import settings
    on_date = on_date or date.today()
    metal = session.get(Metal, job.metal_id) if job.metal_id else None
    sku = job.product_sku
    setting = setting_labour(session, job)
    line = None if live else (_transfer_line(session, job)
                              if job.status == "transferred" else None)
    wip = job.status not in FINISHED
    stock_no = None
    notes: list[str] = []
    if line is not None:
        # Frozen on the transfer: the same inputs give the same figures.
        stones = [mfg_pricing.StoneCost(
            label=st.get("label", ""), pcs=int(st.get("pcs") or 0),
            weight=_dec(st.get("weight")), unit=st.get("unit") or "ct",
            price=_dec(st.get("price")), s_type=st.get("s_type") or "")
            for st in json.loads(line.stones_json or "[]")]
        price = mfg_pricing.price_line(
            net_wt=line.net_wt, title=line.title, fine_rate=line.fine_rate, stones=stones,
            labour_rate=line.labour_rate, labour_weight=line.labour_weight,
            setting_amount=setting, ex_metal_amount=line.ex_metal_amount,
            finding_labour=line.finding_labour, manual_amount=line.manual_amount,
            margin_pct=line.margin_pct, pcs=line.pcs, zero_tag=bool(line.is_repair),
            tag_rule=settings.get_setting(session, mfg_pricing.TAG_DISPLAY_SETTING,
                                          "thousands"))
        gross = _dec(line.gross_wt)
        item = session.scalar(select(StockItem).where(StockItem.line_id == line.id))
        stock_no = item.stock_no if item else None
        from diagold.db.models import MfgTransfer
        t = session.get(MfgTransfer, line.transfer_id)
        basis = (f"Rates frozen on MFG Transfer Vr {t.vr_no} dt {t.vr_date:%d-%m-%Y}"
                 if t else "Rates frozen on the MFG transfer")
        if setting:
            notes.append("Setting labour is the karigar's setting pay on this job; the "
                         "transfer itself carried Setting Amount "
                         f"{_dec(line.setting_amount):,.2f} (28 Sept Q12).")
    else:
        gross_now, net_now = production.last_weights(session, job)
        net = None if not wip else net_now
        price = mfg_pricing.fill_prices(session, job, on_date, setting_amount=setting,
                                        net_wt=net)
        notes += price.notes
        gross = _dec(gross_now)
        basis = (("WIP costing - weights as last weighed, stones issued so far; "
                  if wip else "") + f"rates as on {on_date:%d-%m-%Y}")
    return CostingSheet(
        job_id=job.id, job_no=job.job_no, sku=sku.sku_code if sku else "",
        metal=metal.name if metal else "", colour=job.colour,
        size=production._order_line_size(session, job), pcs=job.pcs, gross_wt=gross,
        price=price, basis=basis, stock_no=stock_no,
        photo=(sku.image_finished or sku.image_design) if sku else "", wip=wip,
        notes=notes)


def _job_date(job: Job) -> date | None:
    return job.completed_on or job.prod_date or (job.created_at.date()
                                                if job.created_at else None)


def register(session: Session, date_from: date, date_to: date, *,
             wip: bool = False, on_date: date | None = None) -> list[dict[str, Any]]:
    """Job Costing register: one row per finished job in the range (by the
    day it finished) - or, with ``wip``, every job still in work (WIP
    costing). Double-click opens the sheet."""
    statuses = WIP if wip else FINISHED
    rows = []
    for job in session.scalars(select(Job).where(Job.status.in_(statuses))
                               .order_by(Job.job_no)):
        d = _job_date(job)
        if not wip and (d is None or d < date_from or d > date_to):
            continue
        sh = costing_sheet(session, job, on_date or date_to)
        if wip and not sh.price.net_wt:
            continue        # nothing weighed yet - nothing to cost
        p = sh.price
        groups = sh.stones_by_group()

        def g(*names: str) -> Decimal:
            return sum((v[2] for k, v in groups.items()
                        if any(n in k.upper() for n in names)), ZERO)
        dia, polki = g("DIA"), g("POLKI")
        rows.append({
            "_job_id": job.id, "job_no": job.job_no, "sku": sh.sku, "metal": sh.metal,
            "col": sh.colour, "size": sh.size, "pcs": sh.pcs, "g_wt": sh.gross_wt,
            "n_wt": p.net_wt, "mt_price": p.metal_rate, "mt_amt": p.metal_amount,
            "st_amt": p.stone_amount, "dia_amt": dia, "polki_amt": polki,
            "cs_amt": p.stone_amount - dia - polki, "setting": sh.setting,
            "std": sh.std_labour, "labour": sh.labour, "total": sh.total,
            "margin_pct": p.margin_pct, "margin": p.margin_amount,
            "price_unit": sh.grand_total, "amount": p.total_value, "tag": p.tag_text,
            "stock": sh.stock_no or "", "status": job.status, "date": _job_date(job),
            "basis": "frozen" if sh.stock_no else ("WIP" if sh.wip else "as on date"),
        })
    return rows


# --------------------------------------------------------------------------
# Excel (2 Oct TR6): numbers, with the arithmetic as live formulas
# --------------------------------------------------------------------------
REGISTER_COLS = [
    ("job_no", "JOBNO"), ("sku", "SKU"), ("metal", "METAL"), ("col", "COL"),
    ("size", "SIZE"), ("pcs", "PCS"), ("g_wt", "G-WT"), ("n_wt", "N-WT"),
    ("mt_price", "MT PRICE"), ("mt_amt", "MT AMT"), ("st_amt", "ST AMT"),
    ("labour", "LABOUR"), ("total", "TOTAL"), ("margin_pct", "MARGIN %"),
    ("margin", "MARGIN"), ("price_unit", "PRICE UNIT"), ("amount", "AMOUNT"),
    ("tag", "TAG PRICE"), ("stock", "STOCK"),
]


def register_xlsx(rows: list[dict[str, Any]], path: str, title: str) -> None:
    """The register as Excel: MT AMT = N-WT x MT PRICE, TOTAL = MT AMT + ST AMT
    + LABOUR, MARGIN = TOTAL x MARGIN %, PRICE UNIT = TOTAL + MARGIN, AMOUNT =
    PRICE UNIT x PCS - formulas, so a typed change re-adds; SUM totals."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([h for _k, h in REGISTER_COLS])
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E4E7EC")
    col = {k: get_column_letter(i) for i, (k, _h) in enumerate(REGISTER_COLS, start=1)}
    first = 3
    for n, r in enumerate(rows):
        i = first + n
        vals = []
        for k, _h in REGISTER_COLS:
            v = r.get(k)
            if k == "mt_amt":
                v = f"=ROUND({col['n_wt']}{i}*{col['mt_price']}{i},2)"
            elif k == "total":
                v = f"={col['mt_amt']}{i}+{col['st_amt']}{i}+{col['labour']}{i}"
            elif k == "margin":
                v = f"=ROUND({col['total']}{i}*{col['margin_pct']}{i}/100,2)"
            elif k == "price_unit":
                v = f"={col['total']}{i}+{col['margin']}{i}"
            elif k == "amount":
                v = f"={col['price_unit']}{i}*{col['pcs']}{i}"
            elif k == "tag":
                v = float(v) if v not in (None, "") and str(v).isdigit() else v
            elif isinstance(v, Decimal):
                v = float(v)
            vals.append(v)
        ws.append(vals)
        for k in ("g_wt", "n_wt"):
            ws[f"{col[k]}{i}"].number_format = "#,##0.000"
        for k in ("mt_price", "mt_amt", "st_amt", "labour", "total", "margin",
                  "price_unit", "amount"):
            ws[f"{col[k]}{i}"].number_format = "#,##0.00"
    last = first + len(rows) - 1
    total_row = ["TOTAL"] + [""] * (len(REGISTER_COLS) - 1)
    ws.append(total_row)
    t = ws.max_row
    if rows:
        for k in ("pcs", "g_wt", "n_wt", "mt_amt", "st_amt", "labour", "total", "margin",
                  "amount"):
            ws[f"{col[k]}{t}"] = f"=SUM({col[k]}{first}:{col[k]}{last})"
            ws[f"{col[k]}{t}"].number_format = "#,##0.000" if k in ("g_wt", "n_wt") \
                else ("0" if k == "pcs" else "#,##0.00")
    for c in ws[t]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9DDE5")
    ws.freeze_panes = "A3"
    for i, (_k, h) in enumerate(REGISTER_COLS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, len(h) + 4)
    wb.save(path)


def sheet_xlsx(sh: CostingSheet, path: str) -> None:
    """One Job Costing Sheet as Excel, every amount a formula of its inputs."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    p = sh.price
    wb = Workbook()
    ws = wb.active
    ws.title = f"Job {sh.job_no}"
    ws.append([f"Job Costing Sheet — Job {sh.job_no}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"{sh.sku} · {sh.metal} {sh.colour} · {sh.pcs} pc · size {sh.size}"])
    ws.append([sh.basis])
    ws.append([])
    ws.append(["Item", "Pcs", "Weight / Ct", "Rate", "Amount"])
    for c in ws[5]:
        c.font = Font(bold=True)
    ws.append(["Metal (N-Wt)", "", float(p.net_wt), float(p.metal_rate), "=ROUND(C6*D6,2)"])
    ws.append(["G-Wt", "", float(sh.gross_wt), "", ""])
    first = ws.max_row + 1
    for s in p.stones:
        qty_col = "B" if s.unit == "pcs" else "C"
        ws.append([f"Stone {s.label}", s.pcs, float(s.weight), float(s.price), ""])
        r = ws.max_row
        ws[f"E{r}"] = f"=ROUND({qty_col}{r}*D{r},2)"
    last = ws.max_row
    ws.append(["Stone total", f"=SUM(B{first}:B{last})" if p.stones else 0,
               f"=SUM(C{first}:C{last})" if p.stones else 0, "",
               f"=SUM(E{first}:E{last})" if p.stones else 0])
    st = ws.max_row
    ws.append(["Setting", "", "", "", float(sh.setting)])
    se = ws.max_row
    ws.append(["STD labour", "", float(p.labour_weight), float(p.labour_rate),
               f"=ROUND(C{se + 1}*D{se + 1},2)"])
    lab = ws.max_row
    extra = float(p.finding_labour + p.ex_metal_amount + p.manual_amount)
    ws.append(["Finding / Ex metal / Manual", "", "", "", extra])
    ex = ws.max_row
    ws.append(["Total", "", "", "", f"=E6+E{st}+E{se}+E{lab}+E{ex}"])
    tot = ws.max_row
    ws.append(["Margin %", "", "", float(p.margin_pct), f"=ROUND(E{tot}*D{tot + 1}/100,2)"])
    mar = ws.max_row
    ws.append(["Grand Total", "", "", "", f"=E{tot}+E{mar}"])
    gt = ws.max_row
    ws.append(["Tag price (Grand Total / 1000, provisional - Q3)", "", "", "",
               f"=INT(E{gt}/1000)"])
    tg = ws.max_row
    ws.append(["Price per gm (tag / G-Wt)", "", "", "", "=ROUND(E%d/C7,2)" % tg])
    for r in (tot, gt):
        for c in ws[r]:
            c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=6, max_row=ws.max_row):
        for c in row[2:5]:
            c.number_format = "#,##0.000" if c.column == 3 else "#,##0.00"
    ws.column_dimensions["A"].width = 46
    for c in "BCDE":
        ws.column_dimensions[c].width = 16
    wb.save(path)
