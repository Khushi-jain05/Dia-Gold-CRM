"""Client-wise price charts - "Prices From Client Chart" (5 Oct §4.7, T-05).

A chart is a price type ("A", "MANNU BHAI") with a default labour per gram
and three rule tables. A rule's match fields are text; a blank field matches
anything. Rules are tried in priority order (1 first) and the first that
matches wins:

* labour  - family, style, SKU, item, metal, colour and a gross-weight slab
            (From G-Wt .. To G-Wt, To 0 = no upper limit) -> sale price per
            gram or per piece. No rule -> the chart's Per Grm Price.
* stone   - family, metal, style, SKU, stone group, SSKU, stone, shape, type,
            quality, size -> price per ct or per piece. No rule -> the
            stone's own (standard) price stays.
* setting - setting type, stone group, size -> price per piece; the piece's
            setting amount = sum of pcs x price over its stones. No rule for
            any stone -> the standard setting amount stays.

Applied on a sale when the client is on a chart and "Prices From Client
Chart" is ticked (default from the setting "Client wise price chart
applicable"). With no chart the standard master price applies.
"""
from __future__ import annotations

from decimal import Decimal
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from diagold.db.models import (Account, FamilyCategory, PriceChart, PriceChartLabour,
                               PriceChartSetting, PriceChartStone, ProductSku, StoneSku)
from diagold.db.models.sku import Item

ZERO = Decimal("0")
APPLICABLE_FLAG = "sale.client_chart"     # "Client wise price chart applicable"
RULE_MODELS = {"labour": PriceChartLabour, "stone": PriceChartStone,
               "setting": PriceChartSetting}
RULE_ATTR = {"labour": "labour_rules", "stone": "stone_rules", "setting": "setting_rules"}


class PriceChartError(ValueError):
    pass


def _dec(v: Any) -> Decimal:
    if v in (None, ""):
        return ZERO
    return v if isinstance(v, Decimal) else Decimal(str(v))


def _norm(v: Any) -> str:
    return " ".join(str(v or "").split()).casefold()


def _match(rule_value: str, actual: Any) -> bool:
    """A blank rule field matches anything; otherwise equal, ignoring case."""
    return not (rule_value or "").strip() or _norm(rule_value) == _norm(actual)


def applicable(session: Session) -> bool:
    from diagold.services import settings
    return settings.flag(APPLICABLE_FLAG, False, session=session)


def chart_for(session: Session, account_id: int | None) -> PriceChart | None:
    a = session.get(Account, account_id) if account_id else None
    return session.get(PriceChart, a.price_chart_id) if a and a.price_chart_id else None


# --------------------------------------------------------------------------
# What a piece / a stone is, in the words the rules use
# --------------------------------------------------------------------------
def piece_keys(session: Session, info: dict[str, Any]) -> dict[str, Any]:
    sku = session.get(ProductSku, info.get("product_sku_id")) \
        if info.get("product_sku_id") else None
    fam = session.get(FamilyCategory, sku.family_id) if sku and sku.family_id else None
    item = session.get(Item, sku.item_id) if sku and sku.item_id else None
    return {"family": fam.name if fam else "", "style": sku.style if sku else "",
            "sku": sku.sku_code if sku else (info.get("sku") or ""),
            "item": item.name if item else "", "metal": info.get("metal") or "",
            "colour": info.get("colour") or "", "gross": _dec(info.get("gross_wt"))}


def stone_keys(session: Session, st: dict[str, Any]) -> dict[str, Any]:
    """A stone line of the piece ("POLKI 12-14", S Type Polki, size 12-14)
    with its Stone master fields when the label names a stone SKU."""
    label = (st.get("label") or "").strip()
    size = (st.get("size") or "").strip()
    sku = None
    for code in dict.fromkeys((label, label[: -len(size)].strip() if size and
                               label.endswith(size) else label, label.split(" ")[0])):
        if code:
            sku = session.scalar(select(StoneSku).where(StoneSku.sku_code == code))
            if sku is not None:
                break
    if not size and sku is not None:
        size = sku.size or ""
    if not size and label and " " in label:
        size = label.rsplit(" ", 1)[1]
    code = sku.sku_code if sku else label
    # Legacy charts give SSKU and Size apart ("POLKI" + "12-14"); our stone
    # codes often carry the size ("POLKI 12-14"), so the code without its
    # size is matched too.
    base = code[: -len(size)].strip() if size and code.endswith(size) else code
    return {"group": st.get("s_type") or "", "ssku": code, "ssku_base": base,
            "stone": sku.stone if sku else "", "shape": sku.shape if sku else "",
            "stone_type": sku.stone_type if sku else "",
            "quality": sku.quality if sku else "", "size": size}


# --------------------------------------------------------------------------
# Resolution
# --------------------------------------------------------------------------
def labour_for(chart: PriceChart, keys: dict[str, Any]) -> tuple[Decimal, str, str] | None:
    """(rate, "gm" / "pc", what it came from); None when the chart prices no
    labour (no rule matches and no Per Grm Price) - the standard labour stays."""
    g = keys["gross"]
    for r in chart.labour_rules:
        if not all(_match(getattr(r, f), keys[f]) for f in
                   ("family", "style", "sku", "item", "metal", "colour")):
            continue
        lo, hi = _dec(r.from_gwt), _dec(r.to_gwt)
        if (lo and g < lo) or (hi and g > hi):
            continue
        unit = "pc" if (r.per_unit or "").lower().startswith("p") else "gm"
        return _dec(r.sale_price), unit, f"{chart.name} rule {r.priority}"
    if not _dec(chart.labour_per_gm):
        return None
    return _dec(chart.labour_per_gm), "gm", f"{chart.name} per gram"


def stone_price_for(chart: PriceChart, pkeys: dict[str, Any],
                    skeys: dict[str, Any]) -> tuple[Decimal, str] | None:
    for r in chart.stone_rules:
        if all(_match(getattr(r, f), pkeys[f]) for f in ("family", "metal", "style", "sku")) \
                and _match(r.stone_group, skeys["group"]) \
                and (_match(r.ssku, skeys["ssku"]) or _match(r.ssku, skeys.get("ssku_base"))) \
                and all(_match(getattr(r, f), skeys[f]) for f in
                        ("stone", "shape", "stone_type", "quality", "size")):
            return _dec(r.price), ("pcs" if (r.unit or "").lower().startswith("p") else "ct")
    return None


def setting_price_for(chart: PriceChart, skeys: dict[str, Any]) -> Decimal | None:
    for r in chart.setting_rules:
        if _match(r.setting_type, skeys["group"]) and _match(r.stone_group, skeys["group"]) \
                and _match(r.size, skeys["size"]):
            return _dec(r.price)
    return None


def apply(session: Session, chart: PriceChart, info: dict[str, Any]) -> dict[str, Any]:
    """The piece re-priced from the chart: labour rate (and unit), each
    stone's price, the setting amount. ``_chart_notes`` says what applied."""
    pkeys = piece_keys(session, info)
    lab = labour_for(chart, pkeys)
    notes = [f"Labour {lab[0]} per {lab[1]} ({lab[2]})" if lab else
             "Labour: standard (the chart has no labour price)"]
    stones, setting, set_hit = [], ZERO, False
    for st in info.get("stones") or []:
        st = dict(st)
        skeys = stone_keys(session, st)
        hit = stone_price_for(chart, pkeys, skeys)
        if hit is not None:
            st["price"], st["unit"] = str(hit[0]), hit[1]
            notes.append(f"{st.get('label', '')}: {hit[0]} per {hit[1]}")
        sp = setting_price_for(chart, skeys)
        if sp is not None:
            set_hit = True
            setting += sp * int(st.get("pcs") or 0)
        stones.append(st)
    out = {**info, "stones": stones, "_chart_notes": notes}
    if lab:
        out["labour_rate"], out["labour_per"] = lab[0], lab[1]
    if set_hit:
        out["setting"] = setting
        notes.append(f"Setting {setting}")
    return out


# --------------------------------------------------------------------------
# Editing
# --------------------------------------------------------------------------
def copy_chart(session: Session, chart: PriceChart, name: str) -> PriceChart:
    """Make A Copy: the whole chart, every rule, under a new name."""
    name = (name or "").strip()
    if not name:
        raise PriceChartError("Give the copy a name.")
    if session.scalar(select(PriceChart).where(PriceChart.name == name)):
        raise PriceChartError(f"A chart named {name} already exists.")
    new = PriceChart(name=name, labour_per_gm=chart.labour_per_gm,
                     price_type=chart.price_type, stone_quality=chart.stone_quality,
                     remark=chart.remark)
    for kind, model in RULE_MODELS.items():
        cols = [c.key for c in model.__table__.columns if c.key not in ("id", "chart_id")]
        getattr(new, RULE_ATTR[kind]).extend(
            model(**{c: getattr(r, c) for c in cols}) for r in getattr(chart, RULE_ATTR[kind]))
    session.add(new)
    session.flush()
    return new


def save_rules(session: Session, chart: PriceChart, kind: str,
               rows: list[dict[str, Any]]) -> None:
    """Replace one rule table; the row order is the priority (1 first)."""
    model = RULE_MODELS[kind]
    cols = {c.key for c in model.__table__.columns} - {"id", "chart_id", "priority"}
    rules = getattr(chart, RULE_ATTR[kind])
    rules.clear()
    session.flush()
    for i, row in enumerate(rows, 1):
        rules.append(model(priority=i, **{k: v for k, v in row.items() if k in cols}))
    session.flush()


def clients_on(session: Session, chart: PriceChart) -> list[str]:
    return list(session.scalars(select(Account.name).where(
        Account.price_chart_id == chart.id).order_by(Account.name)))


def delete_chart(session: Session, chart: PriceChart) -> None:
    if clients_on(session, chart):
        raise PriceChartError(
            f"{chart.name} is used by {', '.join(clients_on(session, chart))} - "
            "take them off it in the Account master first.")
    session.delete(chart)


# --------------------------------------------------------------------------
# Excel import / export of a stone chart (the legacy grid's columns)
# --------------------------------------------------------------------------
STONE_XL = {"family": "Family", "metal": "Metal", "style": "Style", "sku": "SKU",
            "stone_group": "StoneGroup", "ssku": "SSKU", "stone": "Stone", "shape": "Shape",
            "stone_type": "Type", "quality": "Quality", "size": "Size", "price": "Price",
            "unit": "Unit", "cost_price": "CostPrice"}


def _cell(v: Any) -> str:
    """A legacy cell as text: "-" or blank = any. Excel turns sizes such as
    6-8 / 12-14 into dates (2026-06-08); they are read back as month-day."""
    from datetime import date, datetime
    if isinstance(v, (datetime, date)):
        return f"{v.month}-{v.day}"
    text = "" if v is None else str(v).strip()
    if isinstance(v, float) and v.is_integer():
        text = str(int(v))
    return "" if text in ("-", "—") else text


def read_stone_excel(path: str) -> tuple[list[dict[str, Any]], list[str]]:
    """Rows of a Client Wise Stone Price sheet, in the sheet's order (that
    order is the priority). Returns (rules, notes)."""
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True, read_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = [_cell(h).casefold() for h in next(it)]
    pos = {k: head.index(v.casefold()) for k, v in STONE_XL.items() if v.casefold() in head}
    if "ssku" not in pos or "price" not in pos:
        raise PriceChartError("The sheet needs at least the SSKU and Price columns.")
    rules, notes, dates = [], [], 0
    from datetime import date
    for r in it:
        if not any(v not in (None, "") for v in r):
            continue
        row: dict[str, Any] = {}
        for k, i in pos.items():
            v = r[i] if i < len(r) else None
            if k == "size" and isinstance(v, date):
                dates += 1
            if k in ("price", "cost_price"):
                row[k] = _dec(v)
            elif k == "unit":
                row[k] = "pc" if _cell(v).lower().startswith("p") else "ct"
            else:
                row[k] = _cell(v)
        if not row.get("ssku"):
            continue
        rules.append(row)
    if dates:
        notes.append(f"{dates} size(s) Excel had turned into dates were read back "
                     "(e.g. 08-Jun → 6-8).")
    return rules, notes


def import_stone_excel(session: Session, chart: PriceChart, path: str) -> tuple[int, list[str]]:
    """Replace the chart's stone rules with the sheet's."""
    rules, notes = read_stone_excel(path)
    save_rules(session, chart, "stone", rules)
    return len(rules), notes


def export_stone_excel(chart: PriceChart, path: str) -> None:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = chart.name[:31]
    ws.append(list(STONE_XL.values()))
    for r in chart.stone_rules:
        ws.append([float(getattr(r, k)) if k in ("price", "cost_price")
                   else ("Cts" if r.unit == "ct" else "Pcs") if k == "unit"
                   else (getattr(r, k) or "-") for k in STONE_XL])
    wb.save(path)


# The client's own stone charts (legacy Client Wise Stone Price), shipped
# with the app: "A" (sent 6 Oct) and MANNU BHAI (sent 9 Oct, labour 1,175 / gm
# as seen on the 5 Oct call).
DEFAULT_CHARTS: tuple[tuple[str, str, str, Decimal], ...] = (
    ("A", "stone_price_A.xlsx", "seed.stone_chart_A", ZERO),
    ("MANNU BHAI", "stone_price_MANNU_BHAI.xlsx", "seed.stone_chart_MANNU_BHAI", Decimal("1175")),
)


def seed_default_charts(session: Session) -> int:
    """Create each shipped chart with its stone prices, once only: a chart
    that already has stone rules, or a later edit, is never touched. Returns
    rows added."""
    from pathlib import Path

    from diagold.services import settings
    added = 0
    for name, file, done_key, per_gm in DEFAULT_CHARTS:
        if settings.get_setting(session, done_key, "") == "1":
            continue
        path = Path(__file__).resolve().parents[1] / "data" / file
        if not path.is_file():
            continue
        settings.set_setting(session, done_key, "1")
        chart = session.scalar(select(PriceChart).where(PriceChart.name == name))
        if chart is not None and chart.stone_rules:
            continue
        if chart is None:
            chart = PriceChart(name=name, price_type="Sale", labour_per_gm=per_gm)
            session.add(chart)
            session.flush()
        n, _notes = import_stone_excel(session, chart, str(path))
        added += n
    return added
