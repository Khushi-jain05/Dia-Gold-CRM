"""Migration data checks (9 Oct Session 5 §12B, T-13).

Run on the legacy data once it is loaded (DB access or the masters export,
C-01) - and on this database any day - to list what will break go-live
balances, as an Excel for the client to correct or to approve "fresh opening
balances" (Data-Request 3e, C-02):

    I1  negative stone / metal balances per location
    I2  pieces in stock with TAG 0
    I3  SKUs (and sold pieces) with no item name
    I4  items / SKUs with no family
    I5  sales with no salesperson
    I6  junk ("1 b.sapp beads") and near-duplicate (citrin / citrine) stone SKUs
    I7  the same piece sold twice without a return between
    I11 sales at 0 or a token amount
    +   repair chains (old stock -> repair job -> new stock), for review
"""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import date
from decimal import Decimal
from difflib import SequenceMatcher
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from diagold.db.models import (Account, Job, MaterialStock, Location, Metal, ProductSku,
                               ReadyVoucher, ReadyVoucherLine, StockItem, StoneSku)
from diagold.db.models.sku import Item

ZERO = Decimal("0")
TOKEN_AMOUNT = Decimal("1000")      # a piece sold for less is "token" (I11: Rs 93, Rs 234)

CHECKS: tuple[tuple[str, str], ...] = (
    ("neg_stone", "I1 Negative stone balance at a location"),
    ("neg_metal", "I1 Negative metal balance at a location"),
    ("tag_zero", "I2 Piece in stock with TAG 0"),
    ("blank_item", "I3 SKU / sold piece with no item"),
    ("blank_family", "I4 Item or SKU with no family"),
    ("no_salesperson", "I5 Sale with no salesperson"),
    ("junk_ssku", "I6 Junk stone SKU code"),
    ("dup_ssku", "I6 Near-duplicate stone SKUs"),
    ("sold_twice", "I7 Same piece sold twice"),
    ("token_sale", "I11 Sale at 0 or a token amount"),
    ("repair_chain", "Repair chain (review)"),
)
LABEL = dict(CHECKS)


def _dec(v: Any) -> Decimal:
    if v in (None, ""):
        return ZERO
    return v if isinstance(v, Decimal) else Decimal(str(v))


def _norm(code: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (code or "").lower())


def run(session: Session, on_date: date | None = None) -> list[dict[str, Any]]:
    """Every finding: {check, key, detail, _job_id}."""
    out: list[dict[str, Any]] = []

    def add(chk: str, key: Any, detail: str, job_id: int | None = None) -> None:
        out.append({"check": LABEL[chk], "key": key, "detail": detail, "_job_id": job_id})

    # I1 - negatives per location
    for r in session.scalars(select(MaterialStock).where(
            (MaterialStock.pcs < 0) | (MaterialStock.weight < 0))):
        loc = session.get(Location, r.location_id)
        if r.material_class == "stone":
            sku = session.get(StoneSku, r.ref_id) if r.ref_id else None
            add("neg_stone", f"{loc.name if loc else '?'} / {(sku.sku_code if sku else r.ref_text)}"
                f"{' ' + r.size if r.size else ''}", f"{r.pcs} pcs / {_dec(r.weight):.3f} ct")
        elif r.material_class == "metal":
            metal = session.get(Metal, r.ref_id) if r.ref_id else None
            add("neg_metal", f"{loc.name if loc else '?'} / {metal.name if metal else r.ref_text}",
                f"{_dec(r.weight):.3f} g")
    # I2
    for it in session.scalars(select(StockItem).where(StockItem.status.in_(
            ("in_stock", "on_approval")))):
        if not _dec(it.tag_price):
            sku = session.get(ProductSku, it.product_sku_id) if it.product_sku_id else None
            add("tag_zero", it.stock_no, f"{sku.sku_code if sku else ''} net {_dec(it.price):,.2f}",
                it.job_id)
    # I3 / I4
    for sku in session.scalars(select(ProductSku)):
        if not sku.item_id:
            add("blank_item", sku.sku_code, "SKU has no Item")
        itm = session.get(Item, sku.item_id) if sku.item_id else None
        if not (sku.family_id or (itm.family_id if itm else None)):
            add("blank_family", sku.sku_code, "no family on the SKU or its item")
    for itm in session.scalars(select(Item)):
        if not (itm.name or "").strip() or itm.name.strip() == "-":
            add("blank_item", itm.code, "Item master row with no name")
    # I5 / I11 / I7
    sold: dict[int, ReadyVoucher] = {}
    q = (select(ReadyVoucherLine, ReadyVoucher).join(ReadyVoucher)
         .where(ReadyVoucher.vr_type.in_(("rs_sale", "rs_sale_return")))
         .order_by(ReadyVoucher.vr_date, ReadyVoucher.id, ReadyVoucherLine.id))
    seen_vr: set[int] = set()
    for l, v in session.execute(q):
        if v.vr_type == "rs_sale_return":
            sold.pop(l.stock_item_id, None)
            continue
        acct = session.get(Account, v.account_id) if v.account_id else None
        who = acct.name if acct else ""
        if v.id not in seen_vr:
            seen_vr.add(v.id)
            if not (v.salesperson or "").strip():
                add("no_salesperson", f"Sale {v.vr_no}", f"{v.vr_date:%d-%m-%Y} {who}")
        if _dec(l.total) < TOKEN_AMOUNT * max(int(l.pcs or 1), 1):
            add("token_sale", f"Sale {v.vr_no}", f"{who}: {int(l.pcs or 0)} pcs for "
                f"{_dec(l.total):,.2f}", l.job_id)
        if l.stock_item_id in sold:
            first = sold[l.stock_item_id]
            item = session.get(StockItem, l.stock_item_id)
            add("sold_twice", item.stock_no if item else l.stock_item_id,
                f"on Sale {first.vr_no} and Sale {v.vr_no}, no return between", l.job_id)
        if l.stock_item_id:
            sold[l.stock_item_id] = v
        sku = session.get(ProductSku, l.product_sku_id) if l.product_sku_id else None
        if sku is not None and not sku.item_id:
            add("blank_item", f"Sale {v.vr_no}", f"{sku.sku_code} sold with no item")
    # I6
    codes = [(s.id, s.sku_code or "") for s in session.scalars(select(StoneSku))]
    for _id, code in codes:
        if re.match(r"^\s*\d+\s+\D", code) or code.strip() in ("", "-"):
            add("junk_ssku", code or "(blank)", "code starts with a number / is blank")
    groups: dict[str, list[str]] = defaultdict(list)
    for _id, code in codes:
        groups[_norm(re.sub(r"^\s*\d+\s+", "", code))[:3]].append(code)
    reported: set[tuple[str, str]] = set()
    for lst in groups.values():
        for i, a in enumerate(lst):
            for b in lst[i + 1:]:
                na, nb = _norm(re.sub(r"^\s*\d+\s+", "", a)), _norm(re.sub(r"^\s*\d+\s+", "", b))
                # Sizes must be the same (EMERALD PEAR 3*4 and 4*5 are two
                # stones); only the words may differ slightly (citrin / citrine).
                same_size = re.findall(r"\d+", na) == re.findall(r"\d+", nb)
                if na and nb and na != nb and same_size and SequenceMatcher(
                        None, na, nb).ratio() >= 0.85:
                    pair = tuple(sorted((a, b)))
                    if pair not in reported:
                        reported.add(pair)
                        add("dup_ssku", f"{pair[0]} / {pair[1]}", "look alike - one stone twice?")
                elif na and na == nb:
                    pair = tuple(sorted((a, b)))
                    if pair not in reported:
                        reported.add(pair)
                        add("dup_ssku", f"{pair[0]} / {pair[1]}", "same code apart from case / "
                            "spaces")
    # repair chains
    from diagold.services import repair
    for job in session.scalars(select(Job).where(Job.repair_item_id.is_not(None),
                                                 Job.status != "cancelled")):
        old = session.get(StockItem, job.repair_item_id)
        if old is not None and not old.repair_of_id:
            add("repair_chain", old.stock_no, repair.chain_text(session, old), job.id)
    return out


def summary(rows: list[dict[str, Any]]) -> list[tuple[str, int]]:
    n: dict[str, int] = defaultdict(int)
    for r in rows:
        n[r["check"]] += 1
    return [(label, n.get(label, 0)) for _k, label in CHECKS]


def to_xlsx(rows: list[dict[str, Any]], path: str) -> None:
    """One Summary sheet (count per check) and one sheet per check."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Dia Gold - migration data checks", date.today().strftime("%d-%m-%Y")])
    ws.append([])
    ws.append(["Check", "Findings", "What to do"])
    for c in ws[3]:
        c.font = Font(bold=True)
    advice = {"I1": "Clean up, or approve fresh opening balances (Data-Request 3e)",
              "I2": "Give each piece its tag price", "I3": "Fill the item on the SKU",
              "I4": "Fill the family", "I5": "Say whether salesperson is required (Q3)",
              "I6": "Merge or delete the stone SKU", "I7": "Check approval -> sale entries",
              "I11": "Confirm the amount", "Re": "For review only"}
    for label, n in summary(rows):
        ws.append([label, n, advice.get(label.split()[0], "")])
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["C"].width = 60
    for _k, label in CHECKS:
        part = [r for r in rows if r["check"] == label]
        if not part:
            continue
        sh = wb.create_sheet(label[:31].replace("/", "-"))
        sh.append(["Key", "Detail"])
        for c in sh[1]:
            c.font = Font(bold=True)
        for r in part:
            sh.append([str(r["key"]), r["detail"]])
        sh.column_dimensions["A"].width = 34
        sh.column_dimensions["B"].width = 70
    wb.save(path)
