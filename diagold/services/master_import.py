"""The client's master lists (legacy Exception ▸ Masters Excel, sent 9 Oct as
"Master sheet.xlsx") and the locations in use (sent as a screenshot).

The sheet has one column per master - Item, Family, Metal, METAL COLOUR,
Stone, Vendors, Clients. Import only ADDS what is missing (matched by name,
ignoring case); nothing already in the system is changed or removed.

What the sheet does not say is filled as follows, and asked (Q94):
* Metal purity - read from the name where it is written ("22KT GOLD 92.25"
  -> 92.25, "14KT CASTING 590" -> 590), else the karat's standard
  (9KT 37.5, 22KT 91.6); silver / others left 0.
* Stone group - name starting "DIA" or with "Diamond" in it -> Diamond, "POLKI" -> Polki,
  everything else Colour Stone.
* Vendors -> Accounts Payable; Clients -> Sundry Debtors.
"""
from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from diagold.db.models import (Account, Colour, FamilyCategory, Location, Metal, StoneGroup,
                               StoneInfo)
from diagold.db.models.sku import Item

DATA = Path(__file__).resolve().parents[1] / "data"
MASTER_FILE = "master_sheet.xlsx"
# The locations in use (screenshot of the legacy location list, 9 Oct).
LEGACY_LOCATIONS = ("DISMENTAL", "GAURANG JI", "HARISH", "MISCELLANEOS", "MUMBAI OFFICE", "PUSH",
                    "puwai", "RAJAT JI", "RAJESH JI", "REPAIR RECEIPT", "SETTING PURIFICATION",
                    "SHARAD JI", "SHARAD JI REPAIR", "SONU JI", "SUNIL JI", "Virtual",
                    "vishal ji", "VISHAL JIDISMENTAL")
COLUMNS = {"item": "Item", "family": "Family", "metal": "Metal", "colour": "METAL COLOUR",
           "stone": "Stone", "vendor": "Vendors", "client": "Clients"}
KARAT_PERCENT = {"9": Decimal("37.5"), "12": Decimal("50"), "14": Decimal("58.5"),
                 "18": Decimal("75"), "22": Decimal("91.6"), "24": Decimal("100")}


def read_sheet(path: str | Path) -> dict[str, list[str]]:
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True, read_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    head = [str(h or "").strip().casefold() for h in next(rows)]
    pos = {k: head.index(v.casefold()) for k, v in COLUMNS.items() if v.casefold() in head}
    out: dict[str, list[str]] = {k: [] for k in COLUMNS}
    for r in rows:
        for k, i in pos.items():
            v = r[i] if i < len(r) else None
            text = " ".join(str(v).split()) if v not in (None, "") else ""
            if text and text not in out[k]:
                out[k].append(text)
    return out


def _code(session: Session, model, name: str, size: int, field: str = "code") -> str:
    """A unique code from the name: letters and digits, upper case."""
    base = re.sub(r"[^A-Z0-9]", "", name.upper())[:size] or "X"
    code, n = base, 1
    col = getattr(model, field)
    while session.scalar(select(func.count()).select_from(model).where(col == code)):
        n += 1
        code = f"{base[:size - len(str(n))]}{n}"
    return code


def _have(session: Session, model, name: str) -> bool:
    return bool(session.scalar(select(func.count()).select_from(model).where(
        func.lower(model.name) == name.lower())))


def metal_purity(name: str) -> tuple[Decimal, str]:
    """(purity as the legacy master writes it, base metal)."""
    up = name.upper()
    if "SILVER" in up:
        return Decimal("0"), "SILVER"
    m = re.search(r"(\d+(?:\.\d+)?)\s*$", up)
    karat = re.match(r"\s*(\d+)\s*KT", up)
    if m and karat and m.group(1) != karat.group(1):
        return Decimal(m.group(1)), "ALLOY" if up.startswith("ALLOY") else "GOLD"
    if karat and karat.group(1) in KARAT_PERCENT:
        return KARAT_PERCENT[karat.group(1)], "GOLD"
    return Decimal("0"), "OTHER"


def stone_group_code(name: str) -> str:
    up = name.upper()
    if "POLKI" in up:
        return "POLKI"
    if up.startswith("DIA") or "DIAMOND" in up:          # not "EMERALD DIA.CUT"
        return "DIA"
    return "CS"


def import_masters(session: Session, path: str | Path,
                   locations: tuple[str, ...] = ()) -> dict[str, int]:
    """Add what is missing. Returns {master: rows added}."""
    lists = read_sheet(path)
    added = {k: 0 for k in ("Item", "Family", "Metal", "Colour", "Stone", "Vendor", "Client",
                            "Location")}
    for name in lists["item"]:
        if not _have(session, Item, name):
            session.add(Item(name=name, code=_code(session, Item, name, 32)))
            session.flush()
            added["Item"] += 1
    for name in lists["family"]:
        if not _have(session, FamilyCategory, name):
            session.add(FamilyCategory(name=name, code=_code(session, FamilyCategory, name, 24)))
            session.flush()
            added["Family"] += 1
    for name in lists["metal"]:
        if not _have(session, Metal, name):
            purity, base = metal_purity(name)
            session.add(Metal(name=name, code=_code(session, Metal, name, 24),
                              purity_fineness=purity, base_metal=base))
            session.flush()
            added["Metal"] += 1
    for name in lists["colour"]:
        if not _have(session, Colour, name):
            session.add(Colour(name=name.title(), code=_code(session, Colour, name, 24)))
            session.flush()
            added["Colour"] += 1
    groups = {g.code.upper(): g for g in session.scalars(select(StoneGroup))}
    for name in lists["stone"]:
        if not _have(session, StoneInfo, name):
            g = groups.get(stone_group_code(name))
            if g is None:
                continue
            session.add(StoneInfo(name=name, code=_code(session, StoneInfo, name, 24),
                                  stone_group_id=g.id))
            session.flush()
            added["Stone"] += 1
    for kind, typ, grp in (("vendor", "Client", "Accounts Payable"),
                           ("client", "Client", "Sundry Debtors")):
        for name in lists[kind]:
            if not _have(session, Account, name):
                session.add(Account(name=name, code=_code(session, Account, name, 24),
                                    account_type=typ, group_name=grp))
                session.flush()
                added["Vendor" if kind == "vendor" else "Client"] += 1
    for name in locations:
        if not _have(session, Location, name):
            session.add(Location(name=name, code=_code(session, Location, name, 24)))
            session.flush()
            added["Location"] += 1
    return added


def seed_master_sheet(session: Session) -> dict[str, int]:
    """Once only, on first start: the client's master sheet and locations."""
    from diagold.services import settings
    key = "seed.master_sheet"
    if settings.get_setting(session, key, "") == "1" or not (DATA / MASTER_FILE).is_file():
        return {}
    settings.set_setting(session, key, "1")
    return import_masters(session, DATA / MASTER_FILE, LEGACY_LOCATIONS)



# --------------------------------------------------------------------------
# All Department Worker sheet (9 Oct): blocks of "WORKER : NAME" | PROCESS
# (Ghat, Setting, Other). OUTSIDE HANDMADE / OUTSIDE SETTING = out-house.
# --------------------------------------------------------------------------
WORKER_FILE = "all_department_workers.xlsx"
OUTSIDE = {"OUTSIDE HANDMADE": "HandMade", "OUTSIDE SETTING": "Setting"}


def read_workers(path: str | Path) -> list[tuple[str, str, bool]]:
    """[(name, department, in_house)] in the sheet's order."""
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True, read_only=True).worksheets[0]
    out: list[tuple[str, str, bool]] = []
    seen: set[str] = set()
    for r in ws.iter_rows(values_only=True):
        for i, v in enumerate(r):
            text = " ".join(str(v).split()) if v not in (None, "") else ""
            if not text.upper().startswith("WORKER :"):
                continue
            name = text.split(":", 1)[1].strip()
            proc = " ".join(str(r[i + 1]).split()) if i + 1 < len(r) and r[i + 1] else ""
            proc = "" if proc in ("-", "—") else proc
            if not name or name.upper() in seen:
                continue
            seen.add(name.upper())
            outside = proc.upper() in OUTSIDE
            out.append((name, OUTSIDE.get(proc.upper(), proc), not outside and bool(proc)))
    return out


def import_workers(session: Session, path: str | Path) -> dict[str, int]:
    """Add every karigar not in the system as a Worker with its department
    and in-house flag; a karigar already there only gets an empty
    department filled. Returns {"added": n, "updated": n}."""
    added = updated = 0
    for name, dept, in_house in read_workers(path):
        acc = session.scalar(select(Account).where(func.lower(Account.name) == name.lower()))
        if acc is None:
            session.add(Account(name=name, code=_code(session, Account, name, 24),
                                account_type="Worker", group_name="Accounts Payable",
                                department=dept, in_house=in_house))
            session.flush()
            added += 1
        elif not acc.department and dept:
            acc.department, acc.in_house = dept, in_house
            updated += 1
    return {"added": added, "updated": updated}


def seed_workers(session: Session) -> dict[str, int]:
    """Once only, on first start: the client's karigar list."""
    from diagold.services import settings
    key = "seed.workers"
    if settings.get_setting(session, key, "") == "1" or not (DATA / WORKER_FILE).is_file():
        return {}
    settings.set_setting(session, key, "1")
    return import_workers(session, DATA / WORKER_FILE)
