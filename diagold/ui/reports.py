"""The shared report engine (18 September, T-01) and the Reports hub.

Every Production-Planning report is a :class:`ReportSpec` - columns, a query
function from ``services/reports.py``, and a few switches - rendered by one
:class:`ReportWidget`: date range in the header, the legacy toolbar (Print,
Search, Set Column, Options, Group, Adv. Filter, Export, Auto Filter), a
filter bar above grouped reports (the client's one complaint: "to see final
setting I scroll to the bottom"), a total row, a row counter, and a query
that runs on a worker thread so the screen never shows "(Not Responding)".
"""
from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Callable

from sqlalchemy import select

from PySide6.QtCore import (
    QAbstractTableModel,
    QDate,
    QModelIndex,
    QObject,
    Qt,
    QThread,
    Signal,
)
from PySide6.QtGui import QColor, QFont, QKeySequence, QShortcut
from PySide6.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QComboBox,
    QDateEdit,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QFormLayout,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QListWidget,
    QListWidgetItem,
    QMessageBox,
    QProgressBar,
    QPushButton,
    QScrollArea,
    QSplitter,
    QStackedWidget,
    QTableView,
    QVBoxLayout,
    QWidget,
)

from diagold.db.session import SessionLocal
from diagold.services import documents, settings
from diagold.ui.crud import auto_fit, fill_width
from diagold.services import inventory as INV
from diagold.services import manufacturing as MF
from diagold.services import production as P
from diagold.services import reports as R
from diagold.services import stone_reports as SR

TINT_KEY = QColor("#FBEEEE")
TINT_MEASURE = QColor("#EAF5EC")
TINT_GROUP = QColor("#E4E7EC")
TINT_TOTAL = QColor("#D9DDE5")
RED = QColor("#B3261E")


# --------------------------------------------------------------------------
# spec
# --------------------------------------------------------------------------
@dataclass
class Col:
    key: str
    label: str
    kind: str = "key"            # "key" (pink) | "measure" (green)
    decimals: int | None = None  # None -> 3 for weights, 2 for values
    group: str = ""              # column group label, e.g. "OPENING"
    total: bool = False          # summed in subtotal / total rows


@dataclass
class ReportSpec:
    key: str
    title: str
    columns: Callable[[date, date], list[Col]]
    query: Callable[..., list[dict]]
    group_by: str | None = None
    filter_column: str | None = None
    footer: Callable[[list[dict]], str] | None = None
    options: dict[str, tuple[str, bool]] = field(default_factory=dict)
    on_activate: Callable[["ReportWidget", dict], None] | None = None
    negative_key: str | None = None
    date_mode: str = "range"         # "range" | "to" (as on one day) | "none"
    # "fy": the financial year; "month": this month so far; "week": today to today + 7, as the legacy
    # analysis reports open (overdays are counted as of the From date).
    date_default: str = "fy"
    note: str = ""
    # Extra actions on the toolbar: (label, shortcut or "", callback(widget)) -
    # e.g. Job Costing's Ctrl+P "Excel Job Costing" (2 Oct T-02).
    actions: list[tuple[str, str, Callable[["ReportWidget"], None]]] = field(
        default_factory=list)
    # Shift+F12 Show Image: rows carry "_photo" (a file path); a photo column
    # is shown on request (5 Oct UX4).
    images: bool = False
    # Ctrl+F1 Show Stone Group: these columns (the DIA / POLKI / CS split)
    # start hidden and the shortcut shows / hides them (5 Oct UX5).
    stone_group_cols: tuple[str, ...] = ()
    # Header pickers passed to the query as keywords: (key, label, loader) where
    # loader(session) -> [(value, text)] - e.g. the account of a ledger.
    pickers: list[tuple[str, str, Callable[[Any], list[tuple[Any, str]]]]] = field(
        default_factory=list)
    # Ctrl+G draws this chart instead of opening Group (the legacy ledger graph):
    # graph(widget) -> None.
    graph: Callable[["ReportWidget"], None] | None = None
    # Columns hidden until their shortcut shows them: (keys, label, columns),
    # e.g. ("F9", "A/c Ids", ("acc_code",)) on the Sales Register.
    toggle_cols: list[tuple[str, str, tuple[str, ...]]] = field(default_factory=list)


def static(cols: list[Col]) -> Callable[[date, date], list[Col]]:
    return lambda _a, _b: cols


def _is_day_key(key: str) -> bool:
    return key.count("-") == 2 and key[:2].isdigit()


# --------------------------------------------------------------------------
# formatting
# --------------------------------------------------------------------------
def fmt(value: Any, col: Col | None = None) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return "Y" if value else "N"
    if isinstance(value, (date, datetime)):
        return value.strftime("%d-%m-%y")
    if isinstance(value, Decimal):
        places = col.decimals if col and col.decimals is not None else (
            3 if col and ("wt" in col.key or "weight" in col.key) else 2)
        return f"{value:,.{places}f}"
    if isinstance(value, float):
        places = col.decimals if col and col.decimals is not None else 2
        return f"{value:,.{places}f}"
    if isinstance(value, int) and col is not None and col.kind == "measure":
        return f"{value:,}"
    return str(value)


def _is_num(v: Any) -> bool:
    return isinstance(v, (int, float, Decimal)) and not isinstance(v, bool)


# --------------------------------------------------------------------------
# model
# --------------------------------------------------------------------------
class ReportModel(QAbstractTableModel):
    """Rows are (kind, dict): kind = row | group | subtotal | total."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.cols: list[Col] = []
        self.rows: list[tuple[str, dict]] = []
        self.negative_key: str | None = None
        self.row_numbers = False
        self._pix: dict[str, Any] = {}

    def load(self, cols: list[Col], rows: list[tuple[str, dict]]) -> None:
        self.beginResetModel()
        self.cols, self.rows = cols, rows
        self.endResetModel()

    def rowCount(self, parent=QModelIndex()) -> int:  # noqa: N802
        return 0 if parent.isValid() else len(self.rows)

    def columnCount(self, parent=QModelIndex()) -> int:  # noqa: N802
        return 0 if parent.isValid() else len(self.cols) + (1 if self.row_numbers else 0)

    def _col(self, c: int) -> Col | None:
        if self.row_numbers:
            c -= 1
        return self.cols[c] if 0 <= c < len(self.cols) else None

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):  # noqa: N802
        if orientation != Qt.Orientation.Horizontal or role != Qt.ItemDataRole.DisplayRole:
            return None
        if self.row_numbers and section == 0:
            return "#"
        col = self._col(section)
        if col is None:
            return None
        return f"{col.group} · {col.label}" if col.group else col.label

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        kind, row = self.rows[index.row()]
        col = self._col(index.column())
        first = 1 if self.row_numbers else 0
        if col is not None and col.key == "_photo":
            if role == Qt.ItemDataRole.DecorationRole and kind == "row":
                path = row.get("_photo") or ""
                if path not in self._pix:
                    from PySide6.QtGui import QPixmap
                    pm = QPixmap(path) if path else QPixmap()
                    self._pix[path] = pm.scaled(56, 56, Qt.AspectRatioMode.KeepAspectRatio,
                                                Qt.TransformationMode.SmoothTransformation) \
                        if not pm.isNull() else None
                return self._pix[path]
            return "" if role == Qt.ItemDataRole.DisplayRole else None
        if role == Qt.ItemDataRole.DisplayRole:
            if self.row_numbers and index.column() == 0:
                return str(row.get("_n", "")) if kind == "row" else ""
            if col is None:
                return ""
            if kind == "group":
                return row.get("_label", "") if index.column() == first else ""
            return fmt(row.get(col.key), col)
        if role == Qt.ItemDataRole.TextAlignmentRole:
            v = row.get(col.key) if col else None
            if _is_num(v) or (col and col.kind == "measure"):
                return int(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            return int(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter)
        if role == Qt.ItemDataRole.BackgroundRole:
            if kind == "group":
                return TINT_GROUP
            if kind in ("subtotal", "total"):
                return TINT_TOTAL
            if col is None:
                return None
            return TINT_MEASURE if col.kind == "measure" else TINT_KEY
        if role == Qt.ItemDataRole.FontRole and kind in ("group", "subtotal", "total"):
            f = QFont()
            f.setBold(True)
            return f
        if role == Qt.ItemDataRole.ForegroundRole and kind == "row":
            if self.negative_key and row.get(self.negative_key):
                return RED
            if col is not None:
                v = row.get(col.key)
                if _is_num(v) and v < 0:
                    return RED
        return None

    def row_at(self, r: int) -> tuple[str, dict] | None:
        return self.rows[r] if 0 <= r < len(self.rows) else None


# --------------------------------------------------------------------------
# worker
# --------------------------------------------------------------------------
# Every report thread lives here until it finishes, owned by no widget: a
# report closed (tab, drill-down dialog) while still loading must never take
# its running thread with it - Qt aborts the whole app on that ("QThread:
# Destroyed while thread is still running"; 6 Oct: the app closed on a click).
_LIVE_THREADS: set = set()


def _forget_thread(thread: "QThread", worker: "QObject") -> None:
    _LIVE_THREADS.discard((thread, worker))
    thread.deleteLater()
    worker.deleteLater()


def wait_for_reports(timeout_ms: int = 5000) -> None:
    """At exit: let running report threads end before Qt tears down."""
    for thread, _w in list(_LIVE_THREADS):
        thread.quit()
        thread.wait(timeout_ms)


class _Worker(QObject):
    done = Signal(object)
    failed = Signal(str)
    ended = Signal()          # always fires, so the thread's loop always quits

    def __init__(self, fn: Callable[..., list[dict]], args: tuple, kwargs: dict):
        super().__init__()
        self.fn, self.args, self.kwargs = fn, args, kwargs
        self.cancelled = False

    def run(self) -> None:
        try:
            try:
                with SessionLocal() as s:
                    rows = self.fn(s, *self.args, **self.kwargs)
            except Exception as exc:  # noqa: BLE001 - shown to the user
                if not self.cancelled:
                    self.failed.emit(str(exc))
                return
            if not self.cancelled:
                self.done.emit(rows)
        finally:
            self.ended.emit()


# --------------------------------------------------------------------------
# dialogs
# --------------------------------------------------------------------------
class _ColumnsDialog(QDialog):
    def __init__(self, cols: list[Col], visible: list[str], parent=None):
        super().__init__(parent)
        self.setWindowTitle("Set Column")
        self.setMinimumSize(380, 460)
        lay = QVBoxLayout(self)
        lay.addWidget(QLabel("Tick the columns to show; move them to change the order."))
        self.list = QListWidget()
        keys = {c.key for c in cols}
        order = [k for k in visible if k in keys]
        order += [c.key for c in cols if c.key not in order]
        by_key = {c.key: c for c in cols}
        for k in order:
            c = by_key[k]
            it = QListWidgetItem(f"{c.group} · {c.label}" if c.group else c.label)
            it.setData(Qt.ItemDataRole.UserRole, k)
            it.setFlags(it.flags() | Qt.ItemFlag.ItemIsUserCheckable)
            it.setCheckState(Qt.CheckState.Checked if k in visible else Qt.CheckState.Unchecked)
            self.list.addItem(it)
        lay.addWidget(self.list, 1)
        bar = QHBoxLayout()
        for label, d in (("Up", -1), ("Down", 1)):
            b = QPushButton(label)
            b.clicked.connect(lambda _=False, d=d: self._move(d))
            bar.addWidget(b)
        b_all = QPushButton("All")
        b_all.clicked.connect(lambda: self._set_all(True))
        b_none = QPushButton("None")
        b_none.clicked.connect(lambda: self._set_all(False))
        bar.addWidget(b_all)
        bar.addWidget(b_none)
        bar.addStretch(1)
        lay.addLayout(bar)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok
                                   | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        lay.addWidget(buttons)

    def _move(self, d: int) -> None:
        r = self.list.currentRow()
        if r < 0 or not (0 <= r + d < self.list.count()):
            return
        it = self.list.takeItem(r)
        self.list.insertItem(r + d, it)
        self.list.setCurrentRow(r + d)

    def _set_all(self, on: bool) -> None:
        for i in range(self.list.count()):
            self.list.item(i).setCheckState(Qt.CheckState.Checked if on else Qt.CheckState.Unchecked)

    def visible(self) -> list[str]:
        return [self.list.item(i).data(Qt.ItemDataRole.UserRole)
                for i in range(self.list.count())
                if self.list.item(i).checkState() == Qt.CheckState.Checked]


class _AutoFilterDialog(QDialog):
    """Per-column: distinct values with a count, checkboxes, a text filter."""

    def __init__(self, col: Col, values: dict[str, int], selected: set[str] | None, parent=None):
        super().__init__(parent)
        self.setWindowTitle(f"Auto Filter — {col.label}")
        self.setMinimumSize(360, 460)
        lay = QVBoxLayout(self)
        self.search = QLineEdit()
        self.search.setPlaceholderText("Type to narrow the list…")
        self.search.textChanged.connect(self._narrow)
        lay.addWidget(self.search)
        self.list = QListWidget()
        for text, n in sorted(values.items(), key=lambda kv: kv[0].lower()):
            it = QListWidgetItem(f"{text or '(blank)'}   ({n})")
            it.setData(Qt.ItemDataRole.UserRole, text)
            it.setFlags(it.flags() | Qt.ItemFlag.ItemIsUserCheckable)
            on = selected is None or text in selected
            it.setCheckState(Qt.CheckState.Checked if on else Qt.CheckState.Unchecked)
            self.list.addItem(it)
        lay.addWidget(self.list, 1)
        self.count = QLabel(f"Count: {len(values)}")
        self.count.setObjectName("Muted")
        lay.addWidget(self.count)
        bar = QHBoxLayout()
        b_all = QPushButton("All")
        b_all.clicked.connect(lambda: self._set(True))
        b_none = QPushButton("None")
        b_none.clicked.connect(lambda: self._set(False))
        bar.addWidget(b_all)
        bar.addWidget(b_none)
        bar.addStretch(1)
        lay.addLayout(bar)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok
                                   | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        lay.addWidget(buttons)

    def _narrow(self, text: str) -> None:
        t = text.strip().lower()
        shown = 0
        for i in range(self.list.count()):
            it = self.list.item(i)
            hide = bool(t) and t not in it.text().lower()
            it.setHidden(hide)
            shown += 0 if hide else 1
        self.count.setText(f"Count: {shown}")

    def _set(self, on: bool) -> None:
        for i in range(self.list.count()):
            it = self.list.item(i)
            if not it.isHidden():
                it.setCheckState(Qt.CheckState.Checked if on else Qt.CheckState.Unchecked)

    def selected(self) -> set[str] | None:
        out = {self.list.item(i).data(Qt.ItemDataRole.UserRole)
               for i in range(self.list.count())
               if self.list.item(i).checkState() == Qt.CheckState.Checked}
        return None if len(out) == self.list.count() else out


_OPS = ("contains", "=", "≠", ">", "<", "≥", "≤", "is blank", "not blank")


class _AdvFilterDialog(QDialog):
    def __init__(self, cols: list[Col], current: list[tuple[str, str, str]], parent=None):
        super().__init__(parent)
        self.setWindowTitle("Adv. Filter")
        self.setMinimumWidth(520)
        lay = QVBoxLayout(self)
        lay.addWidget(QLabel("Rows must match every condition."))
        self.rows: list[tuple[QComboBox, QComboBox, QLineEdit]] = []
        for i in range(3):
            h = QHBoxLayout()
            c = QComboBox()
            c.addItem("— none —", None)
            for col in cols:
                c.addItem(col.label, col.key)
            o = QComboBox()
            o.addItems(_OPS)
            v = QLineEdit()
            if i < len(current):
                k, op, val = current[i]
                c.setCurrentIndex(max(c.findData(k), 0))
                o.setCurrentText(op)
                v.setText(val)
            h.addWidget(c, 2)
            h.addWidget(o, 1)
            h.addWidget(v, 2)
            lay.addLayout(h)
            self.rows.append((c, o, v))
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok
                                   | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        lay.addWidget(buttons)

    def conditions(self) -> list[tuple[str, str, str]]:
        return [(c.currentData(), o.currentText(), v.text().strip())
                for c, o, v in self.rows if c.currentData()]


def _match(value: Any, op: str, text: str) -> bool:
    if op == "is blank":
        return value in (None, "")
    if op == "not blank":
        return value not in (None, "")
    if op == "contains":
        return text.lower() in fmt(value).lower()
    if _is_num(value):
        try:
            other: Any = Decimal(text.replace(",", ""))
        except Exception:  # noqa: BLE001
            return False
        v: Any = Decimal(str(value))
    elif isinstance(value, (date, datetime)):
        other = None
        for pattern in ("%d-%m-%y", "%d-%m-%Y", "%Y-%m-%d"):
            try:
                other = datetime.strptime(text, pattern).date()
                break
            except ValueError:
                continue
        if other is None:
            return False
        v = value if isinstance(value, date) else value.date()
    else:
        v, other = fmt(value).lower(), text.lower()
    return {"=": v == other, "≠": v != other, ">": v > other, "<": v < other,
            "≥": v >= other, "≤": v <= other}.get(op, False)


class _OptionsDialog(QDialog):
    def __init__(self, show_total: bool, row_numbers: bool, summary_only: bool, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Options")
        form = QFormLayout(self)
        self.total = QCheckBox("Show total row")
        self.total.setChecked(show_total)
        self.numbers = QCheckBox("Show row numbers")
        self.numbers.setChecked(row_numbers)
        self.summary = QCheckBox("Summary only (group and total rows)")
        self.summary.setChecked(summary_only)
        for w in (self.total, self.numbers, self.summary):
            form.addRow(w)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok
                                   | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        form.addRow(buttons)


# --------------------------------------------------------------------------
# the report widget
# --------------------------------------------------------------------------
REPORT_COL_MAX = 260     # px - a long remark is capped, the user can drag it wider


class ReportWidget(QWidget):
    open_requested = Signal(str)

    def __init__(self, spec: ReportSpec, user=None, parent=None, autorun: bool = True):
        super().__init__(parent)
        self.spec, self.user = spec, user
        self.job_id: int | None = None
        self._raw: list[dict] = []
        self._filtered: list[dict] = []
        self._cols: list[Col] = []
        self._visible: list[str] | None = self._load_columns()
        self._auto: dict[str, set[str]] = {}
        self._adv: list[tuple[str, str, str]] = []
        self._bar_values: set[str] | None = None
        self._bar_boxes: dict[str, QCheckBox] = {}
        self._thread: QThread | None = None
        self._worker: _Worker | None = None
        # Runs that were superseded before they finished: kept alive until
        # their thread ends, because a running QThread must never be deleted.
        self._orphans: list[tuple[QThread, _Worker]] = []
        self._started = datetime.now()
        self.show_total = True
        self.row_numbers = False
        self.summary_only = False

        outer = QVBoxLayout(self)
        outer.setContentsMargins(20, 16, 20, 16)
        outer.setSpacing(8)

        # Three short rows instead of one long one: on a laptop, with the
        # Reports list on the left, a single row of nine buttons plus search
        # and group ran past the right edge of the window.
        head = QHBoxLayout()
        h1 = QLabel(spec.title)
        h1.setObjectName("H1")
        h1.setWordWrap(True)
        head.addWidget(h1, 1)
        fy0, fy1 = R.fy_range()
        if spec.date_default == "week":
            from datetime import timedelta
            fy0, fy1 = date.today(), date.today() + timedelta(days=7)
        elif spec.date_default == "today":
            fy0 = fy1 = date.today()
        elif spec.date_default == "month":
            # The month so far - a settlement is run for one month.
            fy0, fy1 = date.today().replace(day=1), date.today()
        self.d_from = QDateEdit(QDate(fy0.year, fy0.month, fy0.day))
        self.d_to = QDateEdit(QDate(fy1.year, fy1.month, fy1.day))
        for d in (self.d_from, self.d_to):
            d.setCalendarPopup(True)
            d.setDisplayFormat("dd-MM-yyyy")
        if spec.date_mode == "range":
            head.addWidget(QLabel("From"))
            head.addWidget(self.d_from)
            head.addWidget(QLabel("To"))
            head.addWidget(self.d_to)
        elif spec.date_mode == "to":
            # A position on one day (WIP as on …): only the To date, today.
            today = date.today()
            self.d_to.setDate(QDate(today.year, today.month, today.day))
            head.addWidget(QLabel("As on"))
            head.addWidget(self.d_to)
        self.picker_boxes: dict[str, QComboBox] = {}
        if spec.pickers:
            from diagold.db.session import SessionLocal as _SL
            with _SL() as _s:
                for key, label, loader in spec.pickers:
                    head.addWidget(QLabel(label))
                    cb = QComboBox()
                    cb.setEditable(True)
                    cb.setMinimumWidth(200)
                    cb.setInsertPolicy(QComboBox.InsertPolicy.NoInsert)
                    for value, text in loader(_s):
                        cb.addItem(text, value)
                    cb.completer().setFilterMode(Qt.MatchFlag.MatchContains)
                    cb.currentIndexChanged.connect(lambda _i: self.run())
                    self.picker_boxes[key] = cb
                    head.addWidget(cb)
        self.btn_run = QPushButton("Run")
        self.btn_run.setObjectName("Primary")
        self.btn_run.clicked.connect(self.run)
        head.addWidget(self.btn_run)
        outer.addLayout(head)

        row2 = QHBoxLayout()
        row2.setSpacing(6)
        self.search = QLineEdit()
        self.search.setPlaceholderText("Search…")
        self.search.setClearButtonEnabled(True)
        self.search.setMinimumWidth(120)
        self.search.textChanged.connect(self._rebuild)
        row2.addWidget(self.search, 1)
        row2.addWidget(QLabel("Group"))
        self.group = QComboBox()
        self.group.setMinimumWidth(140)
        self.group.currentIndexChanged.connect(lambda _i: self._rebuild())
        row2.addWidget(self.group)
        # A second level (5 Oct T-12: multi-level group with subtotals).
        row2.addWidget(QLabel("then"))
        self.group2 = QComboBox()
        self.group2.setMinimumWidth(120)
        self.group2.currentIndexChanged.connect(lambda _i: self._rebuild())
        row2.addWidget(self.group2)
        outer.addLayout(row2)

        # The buttons wrap onto a second line on a narrow screen instead of
        # pushing the window wider than the laptop (6 Oct: screens cut off).
        from diagold.ui.flow import FlowLayout
        bar = FlowLayout(spacing=6)
        for label, slot in (("Print", self._print), ("Set Column", self._set_columns),
                            ("Options", self._options), ("Adv. Filter", self._adv_filter),
                            ("Export", self._export), ("Auto Filter", self._auto_filter)):
            b = QPushButton(label)
            b.clicked.connect(slot)
            bar.addWidget(b)
        for label, keys, callback in spec.actions:
            b = QPushButton(f"{label}  ({keys})" if keys else label)
            b.clicked.connect(lambda _c=False, cb=callback: cb(self))
            bar.addWidget(b)
            if keys:
                QShortcut(QKeySequence(keys), self, activated=lambda cb=callback: cb(self))
        # F1 Show All, as on every legacy register (2 Oct UX6): every filter
        # off and the whole history in range.
        QShortcut(QKeySequence("F1"), self, activated=self.show_all)
        # Ctrl+G Group (5 Oct UX2): open the group-by list - or, on a ledger,
        # the month-wise graph, as the legacy ledger's Ctrl+G does.
        if spec.graph is not None:
            QShortcut(QKeySequence("Ctrl+G"), self, activated=lambda: spec.graph(self))
            b = QPushButton("Graph  (Ctrl+G)")
            b.clicked.connect(lambda: spec.graph(self))
            bar.addWidget(b)
        else:
            QShortcut(QKeySequence("Ctrl+G"), self,
                      activated=lambda: (self.group.setFocus(), self.group.showPopup()))
        self.show_images = False
        self.show_stone_groups = False
        # Ctrl+E Export, as on the legacy registers.
        QShortcut(QKeySequence("Ctrl+E"), self, activated=self._export)
        self.shown_toggles: set[str] = set()
        for keys, label, tcols in spec.toggle_cols:
            QShortcut(QKeySequence(keys), self,
                      activated=lambda t=tcols: self.toggle_columns(t))
            b = QPushButton(f"{label}  ({keys})")
            b.clicked.connect(lambda _c=False, t=tcols: self.toggle_columns(t))
            bar.addWidget(b)
        if spec.images:
            QShortcut(QKeySequence("Shift+F12"), self, activated=self.toggle_images)
            b = QPushButton("Show Image  (Shift+F12)")
            b.clicked.connect(self.toggle_images)
            bar.addWidget(b)
        if spec.stone_group_cols:
            QShortcut(QKeySequence("Ctrl+F1"), self, activated=self.toggle_stone_groups)
            b = QPushButton("Stone Group  (Ctrl+F1)")
            b.clicked.connect(self.toggle_stone_groups)
            bar.addWidget(b)
        bar.addStretch(1)
        self.option_boxes: dict[str, QCheckBox] = {}
        for key, (label, default) in spec.options.items():
            cb = QCheckBox(label)
            cb.setChecked(default)
            cb.toggled.connect(lambda _c: self.run())
            self.option_boxes[key] = cb
            bar.addWidget(cb)
        outer.addLayout(bar)

        # Filter bar: restrict a grouped report to chosen groups in one click.
        self.filter_box = QWidget()
        fb = QHBoxLayout(self.filter_box)
        fb.setContentsMargins(0, 0, 0, 0)
        fb.setSpacing(6)
        lab = QLabel("Show:")
        lab.setObjectName("Muted")
        fb.addWidget(lab)
        self.filter_scroll = QScrollArea()
        self.filter_scroll.setWidgetResizable(True)
        self.filter_scroll.setFixedHeight(40)
        self.filter_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        self.filter_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self.filter_inner = QWidget()
        self.filter_layout = QHBoxLayout(self.filter_inner)
        self.filter_layout.setContentsMargins(0, 0, 0, 0)
        self.filter_layout.setSpacing(4)
        self.filter_scroll.setWidget(self.filter_inner)
        fb.addWidget(self.filter_scroll, 1)
        self.filter_box.setVisible(bool(spec.filter_column))
        outer.addWidget(self.filter_box)

        self.progress = QProgressBar()
        self.progress.setRange(0, 0)
        self.progress.setTextVisible(False)
        self.progress.setFixedHeight(6)
        self.progress.setVisible(False)
        outer.addWidget(self.progress)

        self.model = ReportModel(self)
        self.model.negative_key = spec.negative_key
        self.view = QTableView()
        self.view.setModel(self.model)
        self.view.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.view.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.view.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.view.verticalHeader().setVisible(False)
        self.view.verticalHeader().setDefaultSectionSize(26)
        self.view.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        # Columns fit their contents and share the spare width, so the grid
        # spans its box even when the last column is hidden (stretching the
        # last section left a blank strip then). A long remark stays capped.
        auto_fit(self.view, REPORT_COL_MAX)
        self.view.doubleClicked.connect(self._activate)
        # F10 does what a double-click does - "Show JB" on the legacy pending
        # grid (28 Sept §4.8, UX4).
        QShortcut(QKeySequence("F10"), self.view,
                  activated=lambda: self._activate(self.view.currentIndex()))
        outer.addWidget(self.view, 1)

        foot = QHBoxLayout()
        self.status = QLabel("")
        self.status.setObjectName("Muted")
        foot.addWidget(self.status, 1)
        self.btn_cancel = QPushButton("Cancel")
        self.btn_cancel.setVisible(False)
        self.btn_cancel.clicked.connect(self._cancel)
        foot.addWidget(self.btn_cancel)
        outer.addLayout(foot)
        if spec.note:
            note = QLabel(spec.note)
            note.setObjectName("Muted")
            note.setWordWrap(True)
            outer.addWidget(note)
        if autorun:
            self.run()

    # -- dates / options ------------------------------------------------
    def dates(self) -> tuple[date, date]:
        a, b = self.d_from.date(), self.d_to.date()
        return date(a.year(), a.month(), a.day()), date(b.year(), b.month(), b.day())

    def set_dates(self, d_from: date, d_to: date) -> None:
        self.d_from.setDate(QDate(d_from.year, d_from.month, d_from.day))
        self.d_to.setDate(QDate(d_to.year, d_to.month, d_to.day))

    def _opts(self) -> dict[str, Any]:
        out: dict[str, Any] = {k: cb.isChecked() for k, cb in self.option_boxes.items()}
        for k, cb in self.picker_boxes.items():
            out[k] = cb.currentData()
        return out

    def set_picker(self, key: str, value: Any) -> None:
        cb = self.picker_boxes.get(key)
        if cb is not None:
            i = cb.findData(value)
            if i >= 0:
                cb.setCurrentIndex(i)

    # -- running ----------------------------------------------------------
    def run(self) -> None:
        if self._thread is not None:
            self._detach()          # let the old run finish on its own, ignored
        d0, d1 = self.dates()
        if d0 > d1:
            self.status.setText("From is after To.")
            return
        self._cols = self.spec.columns(d0, d1)
        self.progress.setVisible(True)
        self.btn_cancel.setVisible(True)
        self.btn_run.setEnabled(False)
        self.status.setText("Running…")
        worker = _Worker(self.spec.query, (d0, d1), self._opts())
        thread = QThread()                       # no parent: see _LIVE_THREADS
        worker.moveToThread(thread)
        _LIVE_THREADS.add((thread, worker))
        thread.started.connect(worker.run)
        # Bound methods: Qt drops these connections if the widget is gone.
        worker.done.connect(self._loaded)
        worker.failed.connect(self._failed)
        worker.ended.connect(thread.quit)
        thread.finished.connect(self._on_thread_finished)
        thread.finished.connect(lambda t=thread, w=worker: _forget_thread(t, w))
        self._worker, self._thread = worker, thread
        self._started = datetime.now()
        thread.start()

    def _detach(self) -> None:
        """Stop listening to the current run without destroying its thread."""
        if self._worker is not None:
            self._worker.cancelled = True
            for sig, slot in ((self._worker.done, self._loaded),
                              (self._worker.failed, self._failed)):
                try:
                    sig.disconnect(slot)
                except (RuntimeError, TypeError):
                    pass
        if self._thread is not None:
            self._orphans.append((self._thread, self._worker))
        self._thread = None
        self._worker = None

    def _on_thread_finished(self) -> None:
        thread = self.sender()
        self._cleanup(thread)

    def _cleanup(self, thread: QThread | None = None) -> None:
        # The thread itself is deleted by _forget_thread, never here.
        if thread is not None and thread is not self._thread:
            self._orphans = [(t, w) for t, w in self._orphans if t is not thread]
            return
        self.progress.setVisible(False)
        self.btn_cancel.setVisible(False)
        self.btn_run.setEnabled(True)
        self._thread = None
        self._worker = None

    def _cancel(self) -> None:
        self._detach()
        self.progress.setVisible(False)
        self.btn_cancel.setVisible(False)
        self.btn_run.setEnabled(True)
        self.status.setText("Cancelled.")

    def _failed(self, message: str) -> None:
        self.status.setText(f"Could not run: {message}")

    def _loaded(self, rows: list[dict]) -> None:
        self._raw = rows
        for n, r in enumerate(rows, start=1):
            r["_n"] = n
        self._auto.clear()
        self._bar_values = None
        self._fill_group_combo()
        self._fill_filter_bar()
        self._rebuild()
        secs = (datetime.now() - self._started).total_seconds()
        self.status.setText(self.status.text() + f"   ·   {secs:.2f} s")

    def wait(self, timeout_ms: int = 10000) -> None:
        """Block until the current run (and any superseded one) finishes."""
        from PySide6.QtWidgets import QApplication
        t0 = datetime.now()
        while (self._thread is not None or self._orphans) and \
                (datetime.now() - t0).total_seconds() * 1000 < timeout_ms:
            QApplication.processEvents()

    def closeEvent(self, event) -> None:  # noqa: N802 - Qt name
        self._detach()          # its thread finishes on its own (_LIVE_THREADS)
        super().closeEvent(event)

    # -- columns / grouping ----------------------------------------------
    def _load_columns(self) -> list[str] | None:
        with SessionLocal() as s:
            raw = settings.get_setting(s, self._columns_key(), "") or \
                settings.get_setting(s, f"report.{self.spec.key}.columns", "")
        try:
            return json.loads(raw) if raw else None
        except ValueError:
            return None

    def _columns_key(self) -> str:
        """Set Column is kept per user (5 Oct T-12)."""
        uid = getattr(self.user, "id", None)
        return f"report.{self.spec.key}.columns" + (f".u{uid}" if uid else "")

    def visible_cols(self) -> list[Col]:
        if not self._visible:
            return list(self._cols)
        by_key = {c.key: c for c in self._cols}
        out = [by_key[k] for k in self._visible if k in by_key]
        # Day columns are born per run and were not around when the order was
        # saved; they still show, after the chosen ones.
        known = set(self._visible)
        out += [c for c in self._cols if c.key not in known and _is_day_key(c.key)]
        return out or list(self._cols)

    def _fill_group_combo(self) -> None:
        self.group.blockSignals(True)
        current = self.group.currentData()
        self.group.clear()
        self.group.addItem("— none —", None)
        for c in self._cols:
            if c.kind == "key" and not _is_day_key(c.key):
                self.group.addItem(c.label, c.key)
        want = current if current is not None else self.spec.group_by
        idx = self.group.findData(want) if want else 0
        self.group.setCurrentIndex(max(idx, 0))
        self.group.blockSignals(False)
        self.group2.blockSignals(True)
        current2 = self.group2.currentData()
        self.group2.clear()
        for i in range(self.group.count()):
            self.group2.addItem(self.group.itemText(i), self.group.itemData(i))
        self.group2.setCurrentIndex(max(self.group2.findData(current2), 0) if current2 else 0)
        self.group2.blockSignals(False)

    def _fill_filter_bar(self) -> None:
        while self.filter_layout.count():
            item = self.filter_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()
        self._bar_boxes = {}
        if not self.spec.filter_column:
            return
        values = sorted({fmt(r.get(self.spec.filter_column)) for r in self._raw},
                        key=lambda v: v.lower())
        b_all = QPushButton("All")
        b_all.setFixedHeight(26)
        b_all.clicked.connect(lambda: self._bar_set(None))
        self.filter_layout.addWidget(b_all)
        for v in values:
            cb = QCheckBox(v or "(blank)")
            cb.setChecked(True)
            cb.toggled.connect(lambda _c: self._bar_changed())
            self._bar_boxes[v] = cb
            self.filter_layout.addWidget(cb)
        self.filter_layout.addStretch(1)

    def _bar_set(self, values: set[str] | None) -> None:
        for v, cb in self._bar_boxes.items():
            cb.blockSignals(True)
            cb.setChecked(values is None or v in values)
            cb.blockSignals(False)
        self._bar_changed()

    def _bar_changed(self) -> None:
        on = {v for v, cb in self._bar_boxes.items() if cb.isChecked()}
        self._bar_values = None if len(on) == len(self._bar_boxes) else on
        self._rebuild()

    def only_group(self, value: str) -> None:
        """Show a single group - what the filter bar does in one click."""
        self._bar_set({value})

    def jump_to_group(self, value: str) -> None:
        for i, (kind, row) in enumerate(self.model.rows):
            if kind == "group" and row.get("_value") == value:
                self.view.scrollTo(self.model.index(i, 0),
                                   QAbstractItemView.ScrollHint.PositionAtTop)
                self.view.selectRow(i)
                return

    # -- filtering + display rows -----------------------------------------
    def _passes(self, row: dict) -> bool:
        term = self.search.text().strip().lower()
        cols = self.visible_cols()
        if term and term not in " ".join(fmt(row.get(c.key), c) for c in cols).lower():
            return False
        for key, allowed in self._auto.items():
            if fmt(row.get(key)) not in allowed:
                return False
        for key, op, text in self._adv:
            if not _match(row.get(key), op, text):
                return False
        if self._bar_values is not None and self.spec.filter_column:
            if fmt(row.get(self.spec.filter_column)) not in self._bar_values:
                return False
        return True

    def _sum_row(self, rows: list[dict], label_key: str | None, label: str) -> dict:
        out: dict[str, Any] = {}
        for c in self._cols:
            if c.total:
                vals = [r.get(c.key) for r in rows if _is_num(r.get(c.key))]
                if vals:
                    out[c.key] = sum(vals)
        if label_key:
            out[label_key] = label
        return out

    def _rebuild(self) -> None:
        cols = self.visible_cols()
        if self.spec.stone_group_cols and not self.show_stone_groups:
            cols = [c for c in cols if c.key not in self.spec.stone_group_cols]
        hidden = {k for _k, _l, ks in self.spec.toggle_cols for k in ks} - self.shown_toggles
        cols = [c for c in cols if c.key not in hidden]
        if self.show_images:
            cols = [Col("_photo", "IMAGE")] + cols
        self._filtered = [r for r in self._raw if self._passes(r)]
        gkeys = [k for k in (self.group.currentData(), self.group2.currentData()) if k]
        gkeys = list(dict.fromkeys(gkeys))
        display: list[tuple[str, dict]] = []
        rows = list(self._filtered)
        first_key = next((c.key for c in cols if c.key != "_photo"), None)
        if gkeys:
            rows.sort(key=lambda r: tuple(fmt(r.get(k)).lower() for k in gkeys)
                      + (r.get("_n", 0),))

            def emit(part: list[dict], level: int) -> None:
                if level == len(gkeys):
                    if not self.summary_only:
                        display.extend(("row", r) for r in part)
                    return
                key = gkeys[level]
                label = next((c.label for c in self._cols if c.key == key), key)
                buckets: dict[str, list[dict]] = {}
                for r in part:
                    buckets.setdefault(fmt(r.get(key)), []).append(r)
                for value, bucket in buckets.items():
                    head = {"_label": ("    " * level) + f"{label.upper()} : {value or '(blank)'}"}
                    if level == 0:
                        head["_value"] = value
                    display.append(("group", head))
                    emit(bucket, level + 1)
                    display.append(("subtotal", self._sum_row(
                        bucket, first_key, ("    " * level) + f"{value} Total ({len(bucket)})")))

            emit(rows, 0)
        elif not self.summary_only:
            display = [("row", r) for r in rows]
        if self.show_total:
            display.append(("total", self._sum_row(rows, first_key, f"TOTAL ({len(rows)})")))
        self.model.row_numbers = self.row_numbers
        self.model.load(cols, display)
        # A long remark must not push every other column off the screen; the
        # user can still drag a column wider.
        fill_width(self.view, REPORT_COL_MAX)
        parts = [f"{len(rows)} row(s)"]
        if len(rows) != len(self._raw):
            parts.append(f"of {len(self._raw)}")
        if self.spec.footer:
            try:
                parts.append(self.spec.footer(rows))
            except Exception:  # noqa: BLE001
                pass
        active = len(self._auto) + len(self._adv) + (1 if self._bar_values is not None else 0)
        if active:
            parts.append(f"{active} filter(s) on")
        self.status.setText("   ·   ".join(parts))

    # -- toolbar actions ------------------------------------------------
    def _set_columns(self) -> None:
        dlg = _ColumnsDialog(self._cols, [c.key for c in self.visible_cols()], self)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        self._visible = dlg.visible()
        with SessionLocal() as s:
            settings.set_setting(s, self._columns_key(), json.dumps(self._visible))
            s.commit()
        self._rebuild()

    def _options(self) -> None:
        dlg = _OptionsDialog(self.show_total, self.row_numbers, self.summary_only, self)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        self.show_total = dlg.total.isChecked()
        self.row_numbers = dlg.numbers.isChecked()
        self.summary_only = dlg.summary.isChecked()
        self._rebuild()

    def _current_col(self) -> Col | None:
        cols = self.visible_cols()
        c = self.view.currentIndex().column() - (1 if self.row_numbers else 0)
        return cols[c] if 0 <= c < len(cols) else (cols[0] if cols else None)

    def distinct_counts(self, key: str) -> dict[str, int]:
        counts: dict[str, int] = {}
        for r in self._raw:
            v = fmt(r.get(key))
            counts[v] = counts.get(v, 0) + 1
        return counts

    def set_auto_filter(self, key: str, values: set[str] | None) -> None:
        if values is None:
            self._auto.pop(key, None)
        else:
            self._auto[key] = values
        self._rebuild()

    def _auto_filter(self) -> None:
        col = self._current_col()
        if col is None:
            return
        dlg = _AutoFilterDialog(col, self.distinct_counts(col.key), self._auto.get(col.key), self)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        self.set_auto_filter(col.key, dlg.selected())

    def set_adv_filter(self, conditions: list[tuple[str, str, str]]) -> None:
        self._adv = conditions
        self._rebuild()

    def _adv_filter(self) -> None:
        dlg = _AdvFilterDialog(self._cols, self._adv, self)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        self.set_adv_filter(dlg.conditions())

    def selected_rows(self, all_if_none: bool = True) -> list[dict]:
        """The data rows picked in the grid (subtotals skipped) - or, with
        nothing picked, every row the filters leave (Catalog / Tag Print)."""
        picked = sorted({i.row() for i in self.view.selectionModel().selectedIndexes()})
        rows = [r for k, r in (self.model.row_at(i) or ("", {}) for i in picked) if k == "row"]
        return rows if rows or not all_if_none else list(self._filtered)

    def toggle_images(self) -> None:
        self.show_images = not self.show_images
        self.view.verticalHeader().setDefaultSectionSize(60 if self.show_images else 26)
        self._rebuild()

    def toggle_columns(self, keys: tuple[str, ...]) -> None:
        for k in keys:
            self.shown_toggles.symmetric_difference_update({k})
        self._rebuild()

    def toggle_stone_groups(self) -> None:
        self.show_stone_groups = not self.show_stone_groups
        self._rebuild()

    def show_all(self) -> None:
        """F1 Show All: search, group, filters off; From the start of records."""
        self.search.clear()
        self._adv = []
        self.group.setCurrentIndex(0)
        self.set_dates(date(2000, 1, 1), max(self.dates()[1], date.today()))
        self.run()

    def _export(self) -> None:
        path, kind = QFileDialog.getSaveFileName(
            self, "Export", f"{self.spec.key}.xlsx", "Excel (*.xlsx);;CSV (*.csv)")
        if not path:
            return
        if path.lower().endswith(".csv") or kind.startswith("CSV"):
            self.export_csv(path)
        else:
            if not path.lower().endswith(".xlsx"):
                path += ".xlsx"
            self.export_xlsx(path)
        QMessageBox.information(self, "Export", f"Saved {path}")

    def export_xlsx(self, path: str) -> int:
        """Excel with numbers as numbers (not text), the header frozen and the
        group / total rows shaded - opens ready to sum (2 Oct TR6 / TR7)."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        cols = self.visible_cols()
        wb = Workbook()
        ws = wb.active
        ws.title = self.spec.title[:31]
        d0, d1 = self.dates()
        ws.append([self.spec.title, f"{d0:%d-%m-%Y} to {d1:%d-%m-%Y}"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([f"{c.group} {c.label}".strip() for c in cols])
        for cell in ws[2]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E4E7EC")
        n = 0
        for kind, row in self.model.rows:
            if kind == "group":
                ws.append([row.get("_label", "")])
                ws.cell(ws.max_row, 1).font = Font(bold=True)
                continue
            vals = []
            for c in cols:
                v = row.get(c.key)
                if isinstance(v, Decimal):
                    v = float(v)
                elif isinstance(v, datetime):
                    v = v.replace(tzinfo=None)
                vals.append(v)
            ws.append(vals)
            r = ws.max_row
            for i, c in enumerate(cols, start=1):
                v = row.get(c.key)
                if isinstance(v, (Decimal, float)):
                    places = c.decimals if c.decimals is not None else (
                        3 if ("wt" in c.key or "weight" in c.key) else 2)
                    ws.cell(r, i).number_format = "#,##0." + "0" * places
                elif isinstance(v, date):
                    ws.cell(r, i).number_format = "dd-mm-yyyy"
            if kind != "row":
                for cell in ws[r]:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor="D9DDE5")
            else:
                n += 1
        ws.freeze_panes = "A3"
        for i, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(c.label) + 4))
        wb.save(path)
        return n

    def export_csv(self, path: str) -> int:
        cols = self.visible_cols()
        n = 0
        with open(path, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow([f"{c.group} {c.label}".strip() for c in cols])
            for kind, row in self.model.rows:
                if kind == "group":
                    w.writerow([row.get("_label", "")])
                    continue
                w.writerow([fmt(row.get(c.key), c) for c in cols])
                n += 1
        return n

    def _print(self) -> None:
        cols = self.visible_cols()
        d0, d1 = self.dates()
        head = "".join(f"<th>{c.group + ' ' if c.group else ''}{c.label}</th>" for c in cols)
        body = []
        for kind, row in self.model.rows:
            if kind == "group":
                body.append(f"<tr style='background:#e4e7ec'><td colspan='{len(cols)}'>"
                            f"<b>{row.get('_label', '')}</b></td></tr>")
                continue
            style = " style='background:#d9dde5;font-weight:bold'" if kind != "row" else ""
            body.append(f"<tr{style}>" + "".join(
                f"<td align='{'right' if _is_num(row.get(c.key)) else 'left'}'>"
                f"{fmt(row.get(c.key), c)}</td>" for c in cols) + "</tr>")
        html = (f"<h2>{self.spec.title}</h2><p>{d0.strftime('%d-%m-%Y')} to "
                f"{d1.strftime('%d-%m-%Y')} · {len(self._filtered)} row(s) · printed "
                f"{datetime.now().strftime('%d-%b-%Y %H:%M')}</p>"
                f"<table border='1' cellspacing='0' cellpadding='3' width='100%'>"
                f"<tr style='background:#eee'>{head}</tr>{''.join(body)}</table>")
        path = documents.PRINT_DIR / f"{self.spec.key}_{datetime.now().strftime('%Y%m%d-%H%M%S')}.pdf"
        documents.to_pdf(html, path)
        QMessageBox.information(self, "Print", f"Saved {path}")

    def _activate(self, index) -> None:
        got = self.model.row_at(index.row())
        if got is None:
            return
        kind, row = got
        if kind == "row" and self.spec.on_activate is not None:
            self.spec.on_activate(self, row)

    def open_job(self, row: dict, key: str = "production_planning.job_history") -> None:
        if row.get("_job_id"):
            self.job_id = row["_job_id"]
            self.open_requested.emit(key)


# --------------------------------------------------------------------------
# report specs
# --------------------------------------------------------------------------
def _day_cols(d0: date, d1: date) -> list[Col]:
    return [Col(R.day_label(d), R.day_label(d), "measure", 0, total=True)
            for d in R._day_keys(d0, d1)]


def _job_analysis_cols(d0: date, d1: date) -> list[Col]:
    return [Col("job_no", "JOB NO"), Col("sku", "SKU"), Col("ord_no", "ORD NO"),
            Col("orddate", "ORDDATE"), Col("refno", "REFNO"), Col("client", "CLIENT"),
            Col("prod_due", "PROD DUE"), Col("delivered", "DELIVERED"),
            Col("overdays", "OVERDAYS", "measure", 0)] + _day_cols(d0, d1) + [
            Col("total", "TOTAL", "measure", 0, total=True)]


def _process_analysis_cols(d0: date, d1: date) -> list[Col]:
    return [Col("process", "PROCESS"), Col("job_no", "JOB NO"), Col("sku", "SKU"),
            Col("ord_no", "ORD NO"), Col("orddate", "ORDDATE"), Col("client", "CLIENT"),
            Col("proc_due", "PROC DUE"), Col("overdays", "OVERDAYS", "measure", 0)] \
        + _day_cols(d0, d1) + [Col("total", "TOTAL", "measure", 0, total=True)]


_LEDGER_COLS = [Col("location", "LOCATION"), Col("group", "STONE GROUP")] + [
    Col(f"{b}_{m}", lab, "measure", dec, group=b.upper(), total=True)
    for b in ("opening", "inward", "outward", "closing")
    for m, lab, dec in (("pcs", "PCS", 0), ("wt", "WEIGHT", 3), ("val", "VALUE", 2))]


def _drill_stock(widget: ReportWidget, row: dict) -> None:
    """The vouchers behind one location x group line."""
    d0, d1 = widget.dates()
    with SessionLocal() as s:
        rows = R.stock_drilldown(s, row["_location_id"], row["group"], d0, d1)
    spec = ReportSpec(
        key="stock_drill", title=f"{row['location']} · {row['group']} — movements",
        columns=static([Col("date", "DATE"), Col("bucket", "PERIOD"), Col("kind", "KIND"),
                        Col("vr", "VR NO"), Col("ref", "REF"), Col("ssku", "SSKU"),
                        Col("size", "SIZE"), Col("job_no", "JOB NO"),
                        Col("pcs", "PCS", "measure", 0, total=True),
                        Col("weight", "WEIGHT", "measure", 3, total=True),
                        Col("value", "VALUE", "measure", 2, total=True)]),
        query=lambda _s, _a, _b, **_k: rows, date_mode="none",
        on_activate=lambda w, r: w.open_job(r))
    from diagold.ui.production import show_in_dialog
    w = ReportWidget(spec, widget.user)
    w.open_requested.connect(lambda key: (setattr(widget, "job_id", w.job_id),
                                          widget.open_requested.emit(key)))
    show_in_dialog(widget, w, spec.title, (1100, 600))


def _drill_metal(widget: ReportWidget, row: dict) -> None:
    """One location's ledger for one metal - the legacy Metal Analysis drill."""
    d0, d1 = widget.dates()
    with SessionLocal() as s:
        rows = INV.location_ledger(s, row["_location_id"], "metal", row["_metal_id"], d0, d1)
    spec = ReportSpec(
        key="metal_drill", title=f"{row['location']} · {row['metal']} — ledger",
        columns=static([Col("date", "DATE"), Col("vrno", "VRNO"), Col("vrtype", "VRTYPE"),
                        Col("particulars", "PARTICULARS"),
                        Col("inward", "INWARD", "measure", 3, total=True),
                        Col("outward", "OUTWARD", "measure", 3, total=True),
                        Col("closing", "CLOSING", "measure", 3)]),
        query=lambda _s, _a, _b, **_k: rows, date_mode="none")
    from diagold.ui.production import show_in_dialog
    show_in_dialog(widget, ReportWidget(spec, widget.user), spec.title, (980, 560))


def _drill_client_metal(widget: ReportWidget, row: dict) -> None:
    """Client Metal O/S drill: that client's metal ledger (5 Oct T-03)."""
    _d0, d1 = widget.dates()
    spec = build_specs()["worker_metal_ledger"]
    with SessionLocal() as s:
        rows = P.worker_metal_ledger(s, date(2000, 1, 1), d1, worker_id=row["_wid"])
    spec = ReportSpec(key="client_metal_drill", title=f"{row['client']} — metal ledger",
                      columns=spec.columns, query=lambda _s, _a, _b, **_k: rows,
                      date_mode="none")
    from diagold.ui.production import show_in_dialog
    show_in_dialog(widget, ReportWidget(spec, widget.user), spec.title, (1200, 600))


def _locations(s) -> list[tuple[Any, str]]:
    from diagold.db.models import Location
    return [(l.id, l.name) for l in s.scalars(select(Location).order_by(Location.name))]


def _register_specs(money, wt, pcs) -> dict[str, ReportSpec]:
    """Loss, Dust and WIP registers (2 Oct T-08) - legacy Inventory ▸ Reports."""
    from diagold.services import registers as RG
    job = [Col("job_no", "JOBNO"), Col("sku", "SKU")]
    order = [Col("ord_no", "ORD NO"), Col("ord_date", "ORD DATE"), Col("ref_no", "REFNO"),
             Col("ccode", "CCODE")]
    pct = lambda k, l: Col(k, l, "measure", 2)                 # noqa: E731
    return {
        "metal_loss_register": ReportSpec(
            key="metal_loss_register", title="Metal Loss Register",
            columns=static([Col("date", "DATE"), Col("vrno", "VRNO"), Col("worker", "WORKER"),
                            Col("process", "PROCESS"), *job, Col("metal", "METAL"),
                            wt("iss_net", "ISSUED N-WT"), wt("rcv_net", "RECEIVED N-WT"),
                            wt("scrap", "SCRAP"), wt("dust", "DUST"), wt("loss", "LOSS WT"),
                            pct("loss_pct", "LOSS %"), pct("alw_pct", "ALW %"),
                            wt("alw_wt", "ALW WT"), wt("excess", "EXCESS"),
                            wt("fine_loss", "FINE LOSS")]),
            query=lambda s, a, b, **_k: RG.metal_loss_register(s, a, b),
            group_by="process", filter_column="process", negative_key="_negative",
            on_activate=lambda w, r: w.open_job(r),
            note="Loss per received step = issued net - received net - scrap - dust, beside "
                 "the allowance. Red: loss above the allowance. Double-click opens the job."),
        "stone_loss_register": ReportSpec(
            key="stone_loss_register", title="Stone Loss Register",
            columns=static([Col("date", "DATE"), Col("vrno", "VRNO"), Col("worker", "WORKER"),
                            Col("process", "PROCESS"), Col("kind", "BROKEN / LOST"),
                            Col("job_no", "JOBNO"), Col("location", "LOCATION"),
                            Col("sku", "SKU"), Col("ssku", "SSKU"), Col("stone", "STONE"),
                            Col("size", "SIZE"), pcs("st_pcs", "PCS"), wt("weight", "WEIGHT"),
                            Col("price", "PRICE", "measure", 2), Col("unit", "UNIT"),
                            money("amount", "AMOUNT"), Col("s_type", "S TYPE"), *order]),
            query=lambda s, a, b, **_k: RG.stone_loss_register(s, a, b),
            group_by="kind", filter_column="kind", on_activate=lambda w, r: w.open_job(r),
            note="Stones broken or lost out of job bags, at the stone's price."),
        "dust_register": ReportSpec(
            key="dust_register", title="Dust Register",
            columns=static([Col("date", "DATE"), Col("vrno", "VRNO"), Col("worker", "WORKER"),
                            Col("process", "PROCESS"), *job, Col("metal", "METAL"),
                            wt("dust", "DUST"), wt("dust_fine", "DUST FINE"),
                            wt("scrap", "SCRAP"), wt("scrap_fine", "SCRAP FINE")]),
            query=lambda s, a, b, **_k: RG.dust_register(s, a, b),
            group_by="worker", filter_column="process", on_activate=lambda w, r: w.open_job(r),
            note="Dust and scrap recorded on the receipts, by karigar and process."),
        "wip_register": ReportSpec(
            key="wip_register", title="WIP Register",
            columns=static([Col("date", "DATE"), Col("vrno", "VRNO"), Col("type", "TYPE"),
                            Col("process", "PROCESS"), Col("particulars", "PARTICULARS"),
                            *job, Col("cref", "CREF"), Col("metal", "METAL"), Col("col", "COL"),
                            Col("size", "SIZE"), pcs("pcs", "PCS"), wt("g_wt", "G-WT"),
                            wt("n_wt", "N-WT"), wt("fine_wt", "FINE-WT"), wt("st_wt", "ST-WT"),
                            wt("ex_wt", "EX-WT"), wt("find_wt", "FIND-WT"),
                            Col("ord_no", "ORDNO"), Col("group", "GROUP"), wt("loss", "LOSS"),
                            money("value", "VALUE")]),
            query=lambda s, a, b, **_k: RG.wip_register(s, a, b),
            group_by="process", filter_column="type", date_mode="to",
            on_activate=lambda w, r: w.open_job(r),
            note="Every job in work as on the To date. WIP = out with the karigar named; PND = "
                 "waiting for its next step. Value = net weight at the day's metal rate + the "
                 "stones in the job at their price. Double-click opens the job."),
        "wip_process_summary": ReportSpec(
            key="wip_process_summary", title="WIP Register — Process Summary",
            columns=static([Col("process", "PROCESS"),
                            Col("pnd_pcs", "PCS", "measure", 1, "PENDING", True),
                            Col("pnd_g", "G-WT", "measure", 3, "PENDING", True),
                            Col("pnd_n", "N-WT", "measure", 3, "PENDING", True),
                            Col("pnd_val", "VALUE", "measure", 2, "PENDING", True),
                            Col("wip_pcs", "PCS", "measure", 1, "WIP", True),
                            Col("wip_g", "G-WT", "measure", 3, "WIP", True),
                            Col("wip_n", "N-WT", "measure", 3, "WIP", True),
                            Col("st_wt", "STONE", "measure", 3, "WIP", True),
                            Col("find_wt", "FINDING", "measure", 3, "WIP", True),
                            Col("mould_wt", "MOULD", "measure", 3, "WIP", True),
                            Col("ex_wt", "EXTRA WT", "measure", 3, "WIP", True),
                            Col("wip_val", "VALUE", "measure", 2, "WIP", True),
                            Col("tot_pcs", "TOT PCS", "measure", 1, "", True),
                            money("tot_val", "TOT VAL"), pcs("tot_wrk", "TOT WRK")]),
            query=lambda s, a, b, **_k: RG.wip_process_summary(s, a, b), date_mode="to",
            note="The WIP Register by process - every process on the master, pending and "
                 "with karigars, and how many karigars are working on each."),
        "wip_stone": ReportSpec(
            key="wip_stone", title="WIP Stone",
            columns=static([Col("location", "LOCATION"), Col("type", "TYPE"), *job,
                            Col("stone_group", "STONE GROUP"), Col("barcode", "BARCODE"),
                            Col("stone", "STONE"), Col("quality", "QUALITY"),
                            Col("ssku", "SSKU"), Col("size", "SIZE"), Col("lot_no", "LOTNO"),
                            pcs("st_pcs", "PCS"), wt("weight", "WEIGHT"),
                            Col("price", "PRICE", "measure", 2), Col("unit", "UNIT"),
                            money("amount", "AMOUNT"), Col("process", "PROCESS"),
                            Col("worker", "WORKER")]),
            query=lambda s, a, b, **_k: RG.wip_stone(s, a, b),
            group_by="stone_group", filter_column="type", date_mode="to",
            on_activate=lambda w, r: w.open_job(r),
            note="Stones in jobs still in work: WIP = out with the karigar named, BAG = still "
                 "in the job bag. Barcode / lot are not recorded on job-bag stones yet."),
    }


def _ST():
    from diagold.services import stock_transfer
    return stock_transfer


def _transfer_cols(money, wt, pcs) -> list[Col]:
    return [Col("date", "DATE"), Col("vrno", "VRNO"), Col("pane", "PANE"),
            Col("location", "LOCATION"), Col("to_location", "TO"), Col("item", "ITEM"),
            Col("size", "SIZE"), pcs("in_pcs", "IN PCS"), wt("in_wt", "IN WT"),
            wt("in_loss_wt", "IN LOSS"), pcs("out_pcs", "OUT PCS"), wt("out_wt", "OUT WT"),
            wt("out_loss_wt", "OUT LOSS"), Col("price", "PRICE", "measure", 2),
            money("amount", "AMOUNT"), Col("contact", "CONTACT"), Col("ref_no", "REFNO")]


def _sale_specs(money, wt, pcs) -> dict[str, ReportSpec]:
    """Sale registers (2 Oct T-09 / T-10) - legacy Sale ▸ Reports."""
    from diagold.services import inventory as INV
    from diagold.services import sales as SL
    piece = [Col("date", "DATE"), Col("vrno", "VRNO"), Col("particulars", "PARTICULARS"),
             Col("ref_no", "REFNO"), Col("barcode", "BARCODE"), Col("sku", "SKU"),
             Col("metal", "METAL"), Col("col", "COL"), pcs("pcs", "PCS"), wt("g_wt", "G-WT"),
             wt("n_wt", "N-WT")]
    vals = [money("dia", "DIA"), money("polki", "POL"), money("cs", "CS"),
            money("labour", "LABOUR"), money("price", "PRICE")]
    stone = [Col("date", "DATE"), Col("vrno", "VRNO"), Col("ref_no", "REFNO"),
             Col("particulars", "PARTICULARS"), Col("location", "LOCATION"),
             Col("ssku", "SSKU"), Col("size", "SIZE"), Col("lot_no", "LOTNO"),
             pcs("pcs", "PCS"), wt("weight", "WEIGHT"), Col("curr", "CURR"),
             Col("price", "PRICE", "measure", 2), Col("unit", "UNIT"), money("amount", "AMOUNT")]
    open_job = lambda w, r: w.open_job(r)                       # noqa: E731
    return {
        "rs_sale_register": ReportSpec(
            key="rs_sale_register", title="Ready Stock Sale Register",
            columns=static([*piece[:2], Col("vrtype", "TYPE"), *piece[2:], money("metal_amount", "METAL AMT"), *vals]),
            query=lambda s, a, b, **_k: SL.ready_register(s, a, b, ("rs_sale", "rs_sale_return")),
            group_by="particulars", filter_column="vrtype", on_activate=open_job,
            note="Every piece sold (RS) or returned (RSR), with its metal, stone and labour."),
        "rs_approval_register": ReportSpec(
            key="rs_approval_register", title="Ready Stock Approval Register",
            columns=static([*piece[:2], Col("vrtype", "TYPE"), *piece[2:], *vals]),
            query=lambda s, a, b, **_k: SL.ready_register(
                s, a, b, ("rs_approval", "rs_approval_return")),
            group_by="particulars", filter_column="vrtype", on_activate=open_job,
            note="Every piece sent on approval (RA) or back from approval (RAR)."),
        "rs_approval_balance": ReportSpec(
            key="rs_approval_balance", title="Ready Stock Approval Balance",
            columns=static([*piece, money("price", "PRICE"), Col("days", "DAYS OUT", "measure", 0)]),
            query=lambda s, a, b, **_k: SL.approval_balance(s, a, b), date_mode="to",
            group_by="particulars", filter_column="particulars", on_activate=open_job,
            note="Pieces out on approval as on the day - with whom and for how long. Balance "
                 "= sent out - returned - sold."),
        "rs_approval_analysis": ReportSpec(
            key="rs_approval_analysis", title="Ready Stock Approval Analysis",
            columns=static([*piece, wt("ret_g", "RET G-WT"), wt("ret_n", "RET N-WT"),
                            wt("ret_fine", "RET FN-WT"), *vals, Col("status", "STATUS"),
                            Col("on", "ON")]),
            query=lambda s, a, b, **_k: SL.approval_analysis(s, a, b),
            group_by="particulars", filter_column="status", on_activate=open_job,
            note="Every piece sent on approval in the period and what became of it: returned "
                 "(weights back), sold, or still out."),
        "metal_sale_register": ReportSpec(
            key="metal_sale_register", title="Metal Sale Register",
            columns=static([Col("date", "DATE"), Col("vrno", "VRNO"), Col("ref_no", "REFNO"),
                            Col("particulars", "PARTICULARS"), Col("location", "LOCATION"),
                            Col("metal", "METAL"), Col("col", "COL"), pcs("pcs", "PCS"),
                            wt("weight", "WEIGHT"), wt("fine", "FINE"), Col("curr", "CURR"),
                            Col("price", "PRICE", "measure", 2), Col("unit", "UNIT"),
                            money("amount", "AMOUNT")]),
            query=lambda s, a, b, **_k: INV.metal_sale_register(s, a, b),
            group_by="particulars", filter_column="metal",
            note="Every Sale ▸ Metal voucher line (e.g. 24KT Gold 10 g @ 15,050)."),
        "stone_sale_register": ReportSpec(
            key="stone_sale_register", title="Stone Sale Register", columns=static(stone),
            query=lambda s, a, b, **_k: INV.stone_register(s, a, b, ("stone_sale",)),
            group_by="particulars", filter_column="ssku",
            note="Every Sale ▸ Stone ▸ Sale line."),
        "stone_approval_register": ReportSpec(
            key="stone_approval_register", title="Stone Approval Register",
            columns=static([stone[0], stone[1], Col("vrtype", "TYPE"), *stone[2:]]),
            query=lambda s, a, b, **_k: INV.stone_register(
                s, a, b, ("stone_approval", "stone_approval_return")),
            group_by="particulars", filter_column="vrtype",
            note="Stones sent on approval (SA) and back (SAR)."),
        "stone_approval_analysis": ReportSpec(
            key="stone_approval_analysis", title="Stone Approval Analysis",
            columns=static([Col("particulars", "PARTICULARS"), Col("ssku", "SSKU"),
                            Col("size", "SIZE"), pcs("out_pcs", "OUT PCS"),
                            wt("out_wt", "OUT WT"), pcs("ret_pcs", "RET PCS"),
                            wt("ret_wt", "RET WT"), pcs("bal_pcs", "BAL PCS"),
                            wt("bal_wt", "BAL WT"), Col("price", "PRICE", "measure", 2),
                            Col("unit", "UNIT"), money("bal_amount", "BAL AMOUNT")]),
            query=lambda s, a, b, **_k: INV.stone_approval_analysis(s, a, b), date_mode="to",
            group_by="particulars", filter_column="particulars",
            note="Per party and stone: sent on approval, returned, and still out."),
        "transfer_register": ReportSpec(
            key="transfer_register", title="Stock Transfer Register",
            columns=static(_transfer_cols(money, wt, pcs)),
            query=lambda s, a, b, **_k: _ST().register(s, a, b),
            group_by="pane", filter_column="pane",
            note="Every Stock Transfer line - pieces out of ready stock, metal and stone in "
                 "and out by location with losses, pieces moved between locations."),
        "melting_list": ReportSpec(
            key="melting_list", title="Melting List",
            columns=static(_transfer_cols(money, wt, pcs)),
            query=lambda s, a, b, **_k: _ST().register(s, a, b, ("ready_out",)),
            group_by="location", filter_column="location",
            note="Ready pieces broken back into metal and stones (Stock Melting)."),
        "rp_register": ReportSpec(
            key="rp_register", title="Ready Items Purchase Register",
            columns=static([*piece[:2], Col("vrtype", "TYPE"), *piece[2:],
                            money("metal_amount", "METAL AMT"), *vals]),
            query=lambda s, a, b, **_k: SL.ready_register(
                s, a, b, ("rp_purchase", "rp_return", "rp_opening")),
            group_by="particulars", filter_column="vrtype",
            note="Pieces bought in ready (RP), returned to the supplier (RPR) or loaded as "
                 "opening stock (OPR), each with its new barcode."),
        "repair_register": ReportSpec(
            key="repair_register", title="Repair Register",
            columns=static([*piece, money("price", "VALUE"), Col("status", "STATUS"),
                            Col("days", "DAYS OUT", "measure", 0)]),
            query=lambda s, a, b, **_k: SL.repair_register(s, a, b),
            group_by="status", filter_column="status", on_activate=open_job,
            note="Every Ready Repair Issue and whether the piece is still out for repair."),
    }


def build_specs() -> dict[str, ReportSpec]:
    from diagold.services.production import order_day_book

    def money(k: str, l: str) -> Col:
        return Col(k, l, "measure", 2, total=True)

    def wt(k: str, l: str) -> Col:
        return Col(k, l, "measure", 3, total=True)

    def pcs(k: str, l: str) -> Col:
        return Col(k, l, "measure", 0, total=True)

    from diagold.ui.job_costing import job_costing_spec

    from diagold.ui import sales_reports as _sales
    from diagold.ui import stock_tools as _stock
    return {**_stock.specs(), **_sales.specs(),
        "job_costing": job_costing_spec(),
        **_register_specs(money, wt, pcs),
        **_sale_specs(money, wt, pcs),
        # -- day books ---------------------------------------------------
        "order_day_book": ReportSpec(
            key="order_day_book", title="Order Day Book",
            columns=static([Col("date", "DATE"), Col("ord_no", "ORD NO"), Col("vr_type", "VR TYPE"),
                            Col("ref_no", "REF NO"), Col("particulars", "PARTICULARS"),
                            Col("family", "Family"), Col("ref", "REF"), Col("metal", "METAL"),
                            Col("loss_pct", "LOSS %", "measure", 2), Col("col", "COL"),
                            Col("size", "SIZE"), pcs("pcs", "PCS"), wt("gwt_pcs", "G-WT/PCS"),
                            Col("del_dt", "DEL-DT"), Col("job_no", "JOB NO"),
                            pcs("job_pcs", "JOB PCS"), Col("prod_dt", "PROD-DT"),
                            Col("priority", "PRIORITY"), Col("remark", "REMARK"),
                            Col("sku", "SKU")]),
            query=lambda s, a, b, **_k: order_day_book(s, a, b)),
        "job_card_day_book": ReportSpec(
            key="job_card_day_book", title="Job Card Day-Book",
            columns=static([Col("date", "DATE"), Col("vrno", "VRNO"), Col("ord_no", "ORD NO"),
                            Col("ref_no", "REF NO"), Col("client", "CLIENT"), Col("sku", "SKU"),
                            Col("design", "DESIGN"), Col("c_ref", "C-REF"), Col("metal", "METAL"),
                            Col("col", "COL"), Col("size", "SIZE"), Col("del_dt", "DEL-DT"),
                            Col("job_no", "JOBNO"), pcs("job_pcs", "JOB PCS"),
                            Col("p_del_dt", "P DEL-DT"), Col("open_loc", "OPEN_LOC"),
                            Col("cancel", "CANCEL")]),
            query=lambda s, a, b, **_k: R.job_card_day_book(s, a, b),
            on_activate=lambda w, r: w.open_job(r)),
        "job_mapping_day_book": ReportSpec(
            key="job_mapping_day_book", title="Job Mapping Day Book",
            columns=static([Col("date", "DATE"), Col("job_no", "JOBNO"), Col("sku", "SKU"),
                            Col("client", "CLIENT"), Col("ord_no", "ORD NO"), Col("route", "ROUTE"),
                            pcs("steps", "STEPS"), Col("first_due", "FIRST DUE"),
                            Col("last_due", "LAST DUE"), Col("status", "STATUS")]),
            query=lambda s, a, b, **_k: R.job_mapping_day_book(s, a, b),
            on_activate=lambda w, r: w.open_job(r, "production_planning.job_mapping")),
        "inv_rtn_stone_day_book": ReportSpec(
            key="inv_rtn_stone_day_book", title="Inv Rtn Stone Day Book",
            columns=static([Col("date", "DATE"), Col("vrno", "VRNO"), Col("location", "LOCATION"),
                            Col("type", "TYPE"), Col("ssku", "SSKU"), Col("size", "SIZE"),
                            Col("stone", "STONE"), Col("job_no", "JOBNO"), pcs("pcs", "PCS"),
                            wt("weight", "WEIGHT"), Col("price_unit", "PRICE UNIT"),
                            money("amount", "AMOUNT")]),
            query=lambda s, a, b, **_k: R.inv_rtn_stone_day_book(s, a, b),
            on_activate=lambda w, r: w.open_job(r, "production_planning.job_card_bag"),
            note="What was returned - or broken - each day, from which job bag, to which "
                 "location. Includes returns posted from the Job Card Bag screen."),
        # -- outstanding -------------------------------------------------
        "inv_rtn_os_stone": ReportSpec(
            key="inv_rtn_os_stone", title="Inv Rtn O/s Stone",
            columns=static([Col("location", "LOCATION"), Col("job_no", "JOBNO"), Col("sku", "SKU"),
                            Col("cref", "CREF"), Col("ssku", "SSKU"), Col("size", "SIZE"),
                            Col("lotno", "LOTNO"), Col("stype", "STYPE"), Col("type", "TYPE"),
                            pcs("pcs", "PCS"), wt("weight", "WEIGHT"),
                            Col("price_unit", "PRICE UNIT"), money("amount", "AMOUNT"),
                            Col("ccode", "CCODE")]),
            query=lambda s, a, b, **_k: R.inv_rtn_os_stone(s, a, b),
            on_activate=lambda w, r: w.open_job(r, "production_planning.job_card_bag"),
            note="Stones still lying in the bags of unfinished jobs. A line drops out when "
                 "the job card is finished or the stones are returned (18 Sept D3)."),
        "job_os_stone": ReportSpec(
            key="job_os_stone", title="Job O/s – Stone (Job Request Outstanding)",
            columns=static([Col("job_no", "JOB NO"), Col("sku", "SKU"), Col("location", "LOCATION"),
                            pcs("job_pcs", "JOB PCS"), Col("ssku", "SSKU"), Col("stone", "STONE"),
                            Col("shape", "SHAPE"), Col("quality", "QUALITY"), Col("size", "SIZE"),
                            pcs("req_pcs", "REQ PCS"), wt("req_wt", "REQ WT"),
                            pcs("iss_pcs", "ISS PCS"), wt("iss_wt", "ISS WT"),
                            pcs("pnd_pcs", "PND PCS"), wt("pnd_wt", "PND WT"),
                            pcs("cl_pcs", "CL PCS"), wt("cl_wt", "CL WT"),
                            Col("ordno", "ORDNO"), Col("orddt", "ORDDT")]),
            query=lambda s, a, b, **_k: R.job_os_stone(s, a, b),
            on_activate=lambda w, r: w.open_job(r, "production_planning.job_card_bag"),
            note="Order-date range. PND = REQ − issued into the bag; CL = requirement of a "
                 "cancelled job. Options ▸ Summary only gives the footer view."),
        "job_os_pct": ReportSpec(
            key="job_os_pct", title="Inv O/S (Job O/s %)",
            columns=static([Col("job_no", "JOBNO"), Col("sku", "SKU"), Col("cref", "CREF"),
                            pcs("st_pcs", "ST PCS"), Col("st_pct", "ST %", "measure", 0),
                            pcs("find_pcs", "FIND PCS"), Col("find_pct", "FIND %", "measure", 0),
                            pcs("mould_pcs", "MOULD PCS"),
                            Col("mould_pct", "MOULD %", "measure", 0),
                            Col("ord_no", "ORD NO"), Col("orddt", "ORDDT"), Col("ref", "REF"),
                            Col("ccode", "CCODE"), Col("deldt", "DELDT")]),
            query=lambda s, a, b, **_k: R.job_os_pct(s, a, b),
            on_activate=lambda w, r: w.open_job(r, "production_planning.job_card_bag"),
            note="The client does not use this one (\"the vendor built it\"); kept last. "
                 "Double-click drills to the Job Card Bag."),
        # -- analysis ----------------------------------------------------
        "job_analysis": ReportSpec(
            key="job_analysis", title="Job Analysis", columns=_job_analysis_cols,
            query=lambda s, a, b, late_only=False, **_k: R.job_analysis(s, a, b, late_only),
            options={"late_only": ("Late deliveries only", False)},
            on_activate=lambda w, r: w.open_job(r), date_default="week",
            note="Overdays are counted as of the From date (today by default). A day column shows 1 when the job "
                 "was open that day. Open = pending / mapped / in progress, cancelled excluded "
                 "- to confirm (18 Sept C-01)."),
        "process_analysis": ReportSpec(
            key="process_analysis", title="Process Analysis (Production Analysis)",
            columns=_process_analysis_cols,
            query=lambda s, a, b, **_k: R.process_analysis(s, a, b),
            group_by="process", filter_column="process",
            on_activate=lambda w, r: w.open_job(r), date_default="week",
            note="Overdays as of the From date (today by default). Grouped by the job's current step: the last step issued but not received, "
                 "else the next step not yet received (to confirm, C-01). Tick a process in "
                 "the Show bar to see only that group - no scrolling to the bottom."),
        "job_card_analysis_stone": ReportSpec(
            key="job_card_analysis_stone",
            title="Job Card Analysis – Stone (Job Inventory – Stone)",
            columns=static(_LEDGER_COLS),
            query=lambda s, a, b, **_k: R.job_card_analysis_stone(s, a, b),
            negative_key="_negative", on_activate=_drill_stock,
            note="CLOSING = OPENING + INWARD − OUTWARD, derived. Returns from job bags count as "
                 "inward, breakage as outward. Negative closings are red: load the opening "
                 "balances (Production Planning ▸ Opening Stone Balances). Double-click a line "
                 "for the vouchers behind it."),
        "job_stock_analysis": ReportSpec(
            key="job_stock_analysis", title="Job Stock Analysis",
            columns=static([Col("ord_dt", "ORD DT"), Col("refno", "REFNO"), Col("party", "PARTY"),
                            Col("job_no", "JOBNO"), Col("sku", "SKU"), Col("stock_dt", "STOCK DT"),
                            Col("days", "DAYS", "measure", 0)]),
            query=lambda s, a, b, **_k: R.job_stock_analysis(s, a, b),
            footer=R.lead_time_footer, on_activate=lambda w, r: w.open_job(r),
            note="DAYS = STOCK DT − ORD DT; blank while the job is open. STOCK DT is the "
                 "receipt of the last route step (MFG transfer once that module exists - Q8)."),
        # -- inventory (28 Sept §4.11-4.14, T-01 / T-10) -----------------
        "metal_analysis": ReportSpec(
            key="metal_analysis", title="Metal Analysis",
            columns=static([Col("location", "LOCATION"), Col("type", "TYPE"),
                            Col("metal", "METAL"), Col("title", "TITLE", "measure", 1),
                            Col("opening", "WEIGHT", "measure", 3, "OPENING", True),
                            Col("inward", "WEIGHT", "measure", 3, "INWARD", True),
                            Col("outward", "WEIGHT", "measure", 3, "OUTWARD", True),
                            Col("closing", "WEIGHT", "measure", 3, "CLOSING", True),
                            Col("fine", "FINE", "measure", 3, "CLOSING", True),
                            Col("base", "BASE")]),
            query=lambda s, a, b, **_k: INV.metal_analysis(s, a, b),
            negative_key="_negative", on_activate=_drill_metal,
            note="Metal held at each location, from the Inventory vouchers. CLOSING = "
                 "OPENING + INWARD − OUTWARD; FINE = weight × title, stored when posted. "
                 "Negative closings are red (28 Sept Q6). Double-click a line for its "
                 "ledger."),
        "client_metal_os": ReportSpec(
            key="client_metal_os", title="Client Metal O/S",
            columns=static([Col("client", "CLIENT"),
                            wt("fine_os", "FINE O/S"), money("amt_os", "AMT O/S"),
                            wt("wip_fine", "WIP FINE"), wt("pnd4stk", "PND4STK"),
                            wt("stk_fine", "STK FINE"), wt("repair_rtn", "RepairRtn"),
                            wt("ret", "Return"), Col("group", "GROUP"),
                            Col("accgroup", "ACCGROUP")]),
            query=lambda s, a, b, **_k: P.client_metal_os(s, b), date_mode="to",
            negative_key="_negative", filter_column="accgroup",
            on_activate=_drill_client_metal,
            note="Each client's balance in fine metal and in money side by side. FINE O/S: "
                 "metal sold and ready pieces sold at their fine, less returns and metal "
                 "received (+ the client owes fine, - we owe). AMT O/S: the ledger balance "
                 "(Dr +). WIP FINE: their jobs in work; PND4STK: finished, waiting for MFG "
                 "Transfer; STK FINE: their pieces in ready stock; RepairRtn: pieces out on a "
                 "repair issue to them; Return: pieces they returned. Double-click a client "
                 "for their metal ledger."),
        "worker_metal_balance": ReportSpec(
            key="worker_metal_balance", title="Worker Balance (Metal)",
            columns=static([Col("group", "GROUP"), Col("worker", "WORKER"), Col("metal", "METAL"),
                            Col("in_wt", "WEIGHT", "measure", 3, "INWARD", True),
                            Col("in_fine", "FINE", "measure", 3, "INWARD", True),
                            Col("out_wt", "WEIGHT", "measure", 3, "OUTWARD", True),
                            Col("out_fine", "FINE", "measure", 3, "OUTWARD", True),
                            wt("loss_wt", "LOSSWT"), wt("loss_fine", "LOSSFINE"),
                            Col("bal_wt", "WEIGHT", "measure", 3, "BALANCE", True),
                            Col("bal_fine", "FINE", "measure", 3, "BALANCE", True),
                            wt("wip_wt", "WIP WT"), wt("wip_fine", "WIP FINE"),
                            Col("process", "PROCESS")]),
            query=lambda s, a, b, **_k: P.party_metal_balance(s, a, b), date_mode="to",
            negative_key="_negative", group_by="group", filter_column="group",
            note="Legacy \"Worker Metal Outstanding\": every worker AND client, per metal - "
                 "inward, outward, loss, the balance owed (after the allowance) and the metal "
                 "still with them on job steps (WIP) with its process. Clients: metal sold, "
                 "ready pieces sold (net weight), metal received. Group by WORKER for subtotals."),
        "account_ledger": ReportSpec(
            key="account_ledger", title="Account Ledger",
            columns=static([Col("ledger", "LEDGER"), Col("date", "DATE"),
                            Col("vrtype", "VRTYPE"), Col("vrno", "VRNO"),
                            Col("narration", "NARRATION"), money("debit", "DEBIT"),
                            money("credit", "CREDIT"), money("balance", "BALANCE"),
                            Col("drcr", "DR/CR")]),
            query=lambda s, a, b, **_k: INV.account_ledger(s, a, b),
            group_by="ledger", filter_column="ledger",
            note="What the purchases post: Dr Purchase A/c, Cr the supplier, for each "
                 "voucher's total. Balance runs per ledger from the start of records."),
        "inv_metal_day_book": ReportSpec(
            key="inv_metal_day_book", title="Metal Day Book",
            columns=static([Col("date", "DATE"), Col("vrtype", "VRTYPE"), Col("vrno", "VRNO"),
                            Col("account", "ACCOUNT"), Col("location", "LOCATION"),
                            Col("item", "METAL"), pcs("pcs", "PCS"), wt("weight", "WEIGHT"),
                            wt("fine", "FINE"), Col("price", "PRICE", "measure", 2),
                            money("amount", "AMOUNT"), wt("wastage", "WASTAGE")]),
            query=lambda s, a, b, **_k: INV.day_book(s, a, b, "metal"),
            group_by="vrtype", filter_column="vrtype",
            note="Every Inventory ▸ Metal voucher line: MP purchase, MI issue, MR receipt."),
        "inv_stone_day_book": ReportSpec(
            key="inv_stone_day_book", title="Stone Day Book",
            columns=static([Col("date", "DATE"), Col("vrtype", "VRTYPE"), Col("vrno", "VRNO"),
                            Col("account", "ACCOUNT"), Col("location", "LOCATION"),
                            Col("item", "SSKU"), Col("size", "SIZE"), pcs("pcs", "PCS"),
                            wt("weight", "WEIGHT"), Col("price", "PRICE", "measure", 2),
                            money("amount", "AMOUNT")]),
            query=lambda s, a, b, **_k: INV.day_book(s, a, b, "stone"),
            group_by="vrtype", filter_column="vrtype",
            note="Every Inventory ▸ Stone voucher line: SP purchase, SI issue, SR receipt. "
                 "Stone issued on job cards is in the Production Planning day books."),
        # -- manufacturing (28 Sept §4.8-4.10) ----------------------------
        "pending_mfg_transfer": ReportSpec(
            key="pending_mfg_transfer", title="Pending for MFG Transfer",
            columns=static([Col("job_no", "JOBNO"), Col("sku", "SKU"), Col("item", "ITEM"),
                            Col("c_ref", "C-REF"), Col("metal", "METAL"), Col("col", "COL"),
                            pcs("pcs", "PCS"), wt("g_wt", "G-WT"), wt("n_wt", "N-WT"),
                            Col("client", "CLIENT"), Col("completed", "FINISHED ON"),
                            money("st_value", "JB ST VALUE")]),
            query=lambda s, a, b, **_k: MF.pending_rows(s), date_mode="none",
            on_activate=lambda w, r: w.open_job(r, "production_planning.job_card_bag"),
            note="Jobs whose last route step has been received and that are not yet in "
                 "ready stock. Transfer them in Manufacturing ▸ MFG Transfer (Show Pending). "
                 "If nothing is pending the list is empty. Double-click (or F10, as in the legacy) "
                 "shows the job's bag."),
        "mfg_transfer_day_book": ReportSpec(
            key="mfg_transfer_day_book", title="MFG Transfer Day Book",
            columns=static([Col("date", "DATE"), Col("vrno", "VRNO"), Col("job_no", "JOBNO"),
                            Col("sku", "SKU"), Col("stock_no", "STOCK NO"), pcs("pcs", "PCS"),
                            wt("g_wt", "G-WT"), wt("n_wt", "N-WT"),
                            money("metal_amt", "METAL AMT"), money("stone_amt", "STONE AMT"),
                            money("labour", "LABOUR"), money("total", "TOTAL"),
                            money("price", "PRICE"), Col("tag", "TAG")]),
            query=lambda s, a, b, **_k: MF.transfer_day_book(s, a, b),
            on_activate=lambda w, r: w.open_job(r),
            note="Every piece put into ready stock in the period, with the prices stored "
                 "at transfer time."),
        "ready_stock": ReportSpec(
            key="ready_stock", title="Ready Stock",
            columns=static([Col("stock_no", "STOCK NO"), Col("sku", "SKU"),
                            Col("job_no", "JOBNO"), Col("client", "CLIENT"),
                            Col("location", "LOCATION"), Col("date", "MF DATE"),
                            Col("vrno", "MF VRNO"), pcs("pcs", "PCS"), wt("g_wt", "G-WT"),
                            wt("n_wt", "N-WT"), money("cost", "COST"), money("price", "PRICE"),
                            Col("tag", "TAG"), Col("printed", "TAG PRINTED")]),
            query=lambda s, a, b, **_k: MF.ready_stock(s, b),
            group_by="location", filter_column="location",
            on_activate=lambda w, r: w.open_job(r), images=True,
            note="Every finished piece in stock on the To date, with its cost, price and "
                 "tag. Double-click opens the job."),
        # -- karigar ledgers (28 Sept R13 / T-04) -----------------------
        "issue_day_book": ReportSpec(
            key="issue_day_book", title="Issue Day Book",
            columns=static([Col("date", "DATE"), Col("time", "TIME"), Col("vrno", "VRNO"),
                            Col("process", "PROCESS"), Col("worker", "WORKER"),
                            Col("job_no", "JOBNO"), Col("sku", "SKU"), pcs("pcs", "PCS"),
                            wt("g_wt", "G-WT"), wt("n_wt", "N-WT"), wt("stone_wt", "STONE")]),
            query=lambda s, a, b, **_k: P.voucher_day_book(s, "issue", a, b),
            group_by="process", filter_column="process",
            on_activate=lambda w, r: w.open_job(r),
            note="Every job step issued to a karigar in the period."),
        "received_day_book": ReportSpec(
            key="received_day_book", title="Received Day Book",
            columns=static([Col("date", "DATE"), Col("time", "TIME"), Col("vrno", "VRNO"),
                            Col("process", "PROCESS"), Col("worker", "WORKER"),
                            Col("job_no", "JOBNO"), Col("sku", "SKU"), pcs("pcs", "PCS"),
                            wt("g_wt", "G-WT"), wt("n_wt", "N-WT"), wt("scrap", "SCRAP"),
                            wt("dust", "DUST"), wt("loss", "LOSS WT"),
                            Col("loss_pct", "LOSS %", "measure", 2),
                            Col("alw_pct", "ALW L %", "measure", 2),
                            money("labour", "LABOUR")]),
            query=lambda s, a, b, **_k: P.voucher_day_book(s, "receive", a, b),
            group_by="process", filter_column="process",
            on_activate=lambda w, r: w.open_job(r),
            note="Every job step received back in the period, with its loss and labour."),
        "worker_stone_ledger": ReportSpec(
            key="worker_stone_ledger", title="Worker Stone Ledger",
            columns=static([Col("group", "GROUP"), Col("worker", "WORKER"), Col("date", "DATE"),
                            Col("vrtype", "VRTYPE"),
                            Col("vrno", "VRNO"), Col("stone", "STONE"), Col("job_no", "JOBNO"),
                            pcs("in_pcs", "IN PCS"), wt("in_wt", "IN WT"),
                            pcs("out_pcs", "OUT PCS"), wt("out_wt", "OUT WT"),
                            pcs("brk_pcs", "BRK/LOS PCS"), wt("brk_wt", "BRK/LOS WT"),
                            Col("bal_pcs", "BAL PCS", "measure", 0),
                            Col("bal_wt", "BAL WT", "measure", 3)]),
            query=lambda s, a, b, **_k: P.worker_stone_ledger(s, a, b),
            group_by="worker", filter_column="worker",
            on_activate=lambda w, r: w.open_job(r),
            note="Stones with each karigar: ISS from a job bag and SI on an Inventory stone "
                 "issue in; BACK to the bag, SET into the piece (on the step's receipt) and SR "
                 "on a stone receipt out."),
        "worker_stone_balance": ReportSpec(
            key="worker_stone_balance", title="Worker Balance (Stone)",
            columns=static([Col("group", "GROUP"), Col("worker", "WORKER"), Col("ssku", "SSKU"),
                            Col("size", "SIZE"),
                            Col("iss_pcs", "PCS", "measure", 0, "ISSUE", True),
                            Col("iss_wt", "WEIGHT", "measure", 3, "ISSUE", True),
                            Col("rtn_pcs", "PCS", "measure", 0, "RETURN", True),
                            Col("rtn_wt", "WEIGHT", "measure", 3, "RETURN", True),
                            Col("set_pcs", "PCS", "measure", 0, "SET", True),
                            Col("set_wt", "WEIGHT", "measure", 3, "SET", True),
                            Col("brk_pcs", "PCS", "measure", 0, "BREAK", True),
                            Col("brk_wt", "WEIGHT", "measure", 3, "BREAK", True),
                            Col("los_pcs", "PCS", "measure", 0, "LOSS", True),
                            Col("los_wt", "WEIGHT", "measure", 3, "LOSS", True),
                            Col("pcs", "PCS", "measure", 0, "CLOSING", True),
                            Col("weight", "WEIGHT", "measure", 3, "CLOSING", True)]),
            query=lambda s, a, b, **_k: P.worker_stone_balance(s, b), date_mode="to",
            group_by="worker", filter_column="group", negative_key="_negative",
            note="Stones each karigar - and each client (stones sold or on approval to "
                 "them) - holds as on the date, by SSKU and size: ISSUE (bag issue, stone "
                 "issue, sale, approval), RETURN (back, receipt, approval return), SET into "
                 "the piece, CLOSING = issue - return - set. BREAK / LOSS: bag stones broken "
                 "or lost with them named - shown, not part of the closing. Negative "
                 "closings in red."),
        "location_stone_balance": ReportSpec(
            key="location_stone_balance", title="Location Wise Stone Balance",
            columns=static([Col("section", ""), Col("line", "PARTICULARS")] + [
                c for g, h in (("dia", "DIA"), ("cs", "CS"), ("pol", "POL"), ("oth", "OTHER"))
                for c in (Col(f"{g}_wt", "CT", "measure", 3, h),
                          Col(f"{g}_pcs", "PCS", "measure", 0, h))]),
            query=lambda s, a, b, location=None, **_k: SR.location_stone_balance(s, location, a, b),
            pickers=[("location", "Location", _locations)], negative_key="_negative",
            note="One location's stones for the period, in carats and pieces per group. Closing "
                 "= Opening + Inward - Outward (loose stones at the location). Balance Details "
                 "add what was issued from here and is still in work: Job Card (in the job's "
                 "bag) and WIP (with karigars / in pieces being made). Negative closings in "
                 "red."),
        "metal_summary": ReportSpec(
            key="metal_summary", title="Metal Summary",
            columns=static([Col("location", "LOCATION"), Col("group", "GROUP"),
                            Col("type", "TYPE"), Col("metal", "METAL"),
                            Col("purity", "PURITY", "measure", 1), Col("where", "WHERE"),
                            wt("net_wt", "NETWT"), wt("fine", "FINE")]),
            query=lambda s, a, b, **_k: SR.metal_summary(s, b), date_mode="to",
            group_by="location", filter_column="where", negative_key="_negative",
            note="Metal by location and stage, as on the date (Metal Analysis stays as it is). "
                 "INV = the location's closing; JC = metal issued against a job not yet with a "
                 "karigar; WIP_PND = jobs waiting for the next process; WIP_@W = job steps out "
                 "with karigars. A job's metal sits with the location it came from."),
        "stone_summary": ReportSpec(
            key="stone_summary", title="Stone Summary",
            columns=static([Col("location", "LOCATION"), Col("group", "GROUP"),
                            Col("stone", "STONE"), Col("where", "WHERE"), pcs("pcs", "PCS"),
                            wt("weight", "WEIGHT"), money("value", "VALUE")]),
            query=lambda s, a, b, **k: SR.stone_summary(s, b, **k), date_mode="to",
            options={"groups_only": ("Groups only (DIAMOND / POLKI / COLOR STONE)", False)},
            group_by="location", filter_column="where", negative_key="_negative",
            note="Every stone by location and where it is: RDY in ready-stock pieces, INV loose "
                 "at the location, JC in a job's bag, WIP with karigars / in pieces being made, "
                 "LOS broken or lost. Tick Groups only to roll up to the three heads."),
        "setting_labour_statement": ReportSpec(
            key="setting_labour_statement", title="Setting Labour Statement",
            columns=static([Col("worker", "WORKER"), Col("month", "MONTH"), Col("date", "DATE"),
                            Col("vrno", "VRNO"), Col("job_no", "JOBNO"), Col("stone", "STONE"),
                            Col("setting_type", "SETTING TYPE"), pcs("issued", "ISSUED"),
                            pcs("back", "BACK"), pcs("set_pcs", "SET PCS"),
                            Col("rate", "RATE / PC", "measure", 2),
                            money("amount", "AMOUNT")]),
            query=lambda s, a, b, **_k: P.setting_labour_statement(s, a, b),
            group_by="worker", filter_column="worker", date_default="month",
            on_activate=lambda w, r: w.open_job(r),
            note="What each karigar is owed for setting: pieces set (issued − back) × the "
                 "setting type's rate on the receive date. Pick the month in From / To; the "
                 "subtotal per karigar is the settlement. Broken pieces that came back are "
                 "unpaid (to confirm, 28 Sept Q4)."),
        "worker_metal_ledger": ReportSpec(
            key="worker_metal_ledger", title="Worker Metal Ledger",
            columns=static([Col("group", "GROUP"), Col("worker", "WORKER"), Col("date", "DATE"),
                            Col("vrno", "VRNO"),
                            Col("vrtype", "VRTYPE"), Col("metal", "METAL"),
                            Col("job_no", "JOBNO"), Col("sku", "SKU"),
                            Col("process", "PROCESS"),
                            Col("in_wt", "WEIGHT", "measure", 3, "INWARD", True),
                            Col("in_fine", "FINEWT", "measure", 3, "INWARD", True),
                            Col("out_wt", "WEIGHT", "measure", 3, "OUTWARD", True),
                            Col("out_fine", "FINEWT", "measure", 3, "OUTWARD", True),
                            Col("loss_wt", "LOSSWT", "measure", 3, "LOSS", True),
                            Col("loss_fine", "LOSSFINE", "measure", 3, "LOSS", True),
                            Col("alw_pct", "ALW L %", "measure", 3, "LOSS"),
                            Col("alw_wt", "ALW L WT", "measure", 3, "LOSS", True),
                            Col("alw_fine", "ALW L FINE", "measure", 3, "LOSS", True),
                            Col("bal_wt", "WEIGHT", "measure", 3, "BALANCE"),
                            Col("bal_fine", "FINE", "measure", 3, "BALANCE")]),
            query=lambda s, a, b, **_k: P.worker_metal_ledger(s, a, b),
            group_by="worker", filter_column="group",
            on_activate=lambda w, r: w.open_job(r),
            note="The karigar's metal account. Inward = net weight issued on a job step; "
                 "outward = net weight received back plus scrap and dust; LOSSWT = what "
                 "was actually lost; ALW L WT = Allow Loss % of the net received. BALANCE "
                 "= opening + inward − outward − allowance, i.e. the karigar owes only loss "
                 "beyond the allowance - switch to 'charged for all loss' in Tools ▸ Option "
                 "(to confirm, 28 Sept Q3). Fine = weight × the job metal's purity."),
        # -- data quality ------------------------------------------------
        "data_quality": ReportSpec(
            key="data_quality", title="Data Quality",
            columns=static([Col("check", "CHECK"), Col("key", "KEY"), Col("detail", "DETAIL")]),
            query=lambda s, a, b, **_k: R.data_quality(s, a, b),
            group_by="check", filter_column="check",
            on_activate=lambda w, r: w.open_job(r),
            note="The migration exception checks from the 18 Sept notes, run on this database. "
                 "They run on the legacy data once server access arrives (S1 C-01)."),
    }


SECTIONS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("Day Books", ("job_card_day_book", "job_mapping_day_book", "inv_rtn_stone_day_book",
                   "order_day_book")),
    ("Outstanding", ("inv_rtn_os_stone", "job_os_stone")),
    ("Analysis", ("job_analysis", "process_analysis", "job_card_analysis_stone",
                  "job_stock_analysis")),
    ("Inventory", ("metal_analysis", "metal_summary", "worker_metal_balance",
                   "inv_metal_day_book",
                   "inv_stone_day_book", "transfer_register", "melting_list",
                   "account_ledger")),
    ("Manufacturing", ("job_costing", "pending_mfg_transfer", "mfg_transfer_day_book",
                       "ready_stock")),
    ("Ready Stock", ("ready_stock", "ready_closing_stock", "sku_status")),
    ("Sale", ("rs_sale_register", "rs_approval_register", "rs_approval_balance",
              "rs_approval_analysis", "metal_sale_register", "stone_sale_register",
              "stone_approval_register", "stone_approval_analysis", "repair_register")),
    ("Sales Analysis", ("sales_register", "sales_return_register", "sales_profit",
                        "purchase_sales", "todays_daybook")),
    ("Purchase", ("rp_register",)),
    ("Registers", ("metal_loss_register", "stone_loss_register", "dust_register",
                   "wip_register", "wip_process_summary", "wip_stone")),
    ("Stone", ("location_stone_balance", "stone_summary", "worker_stone_balance")),
    ("Karigar", ("worker_metal_ledger", "worker_stone_ledger", "worker_stone_balance",
                 "setting_labour_statement", "issue_day_book", "received_day_book")),
    ("Client", ("client_metal_os", "worker_metal_balance")),
    ("Other", ("job_os_pct", "data_quality")),
)
